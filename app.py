"""Aplicação Flask do classificador de violência armada: as rotas e as telas.

ESTE ARQUIVO SÓ TRADUZ HTTP. Ele lê o formulário, chama uma função de domínio e
escolhe o template. A regra de negócio vive em `servico.py`, `classificador.py`,
`corpus.py` e `recebidas.py` — se uma decisão sobre o que o sistema registra
estiver aqui, está no lugar errado.

O CAMINHO DE UMA NOTÍCIA, na ordem em que as telas aparecem:

1. `/recebidas` — a notícia entra com link, data e texto. O BERT olha e a
   classifica numa de três faixas: vermelho (0 votos, pessoa revisa), laranja
   (abaixo do corte, pessoa verifica) ou segue. Sem modelo treinado, tudo cai
   na fila para uma pessoa — e é essa rotulagem que constrói o corpus.
2. `/ocorrencia/nova` — a notícia com fenômeno vira registro. Aqui se pede o
   endereço, que é o que não se sabe antes de ler o texto.
3. `/ocorrencia/<id>` — a LLM aprofunda as classes marcadas: monta a lista de
   vítimas, preenche a ficha de cada uma e as tabelas complementares. O analista
   confere; correção manual prevalece sobre a releitura.
4. `/aprovacao` — ao aprovar, as notícias da ocorrência vão para o corpus de
   treino com os fenômenos confirmados. É aqui que o ciclo fecha.
5. `/modelos` — treina o BERT com o que o corpus acumulou, em subprocesso.

AS OUTRAS TELAS: `/vitimas` consolida as pessoas atingidas e permite busca por
local e data; `/alertas` resolve suspeita de ocorrência duplicada;
`/classificacoes` mostra o que o BERT marcou para conferência, e é dela que o
erro do modelo volta ao corpus; `/corpus` é o repositório de treino;
`/testar` é o playground; `/gestao` edita o catálogo de indicadores;
`/prompts` edita os textos que vão à LLM.

NENHUM PRÉ-PROMPT FICA AQUI. Todo texto que a LLM recebe está em `prompts.py`
(editável em /prompts) ou em `indicadores_fonte.json` (editável em /gestao),
porque instrução escondida em constante de Python é instrução que ninguém
revisa.
"""
import io
import json
import os
from datetime import date

from dotenv import load_dotenv
from flask import (
    Flask, abort, flash, g, redirect, render_template, request, send_file,
    session, url_for
)

import ajuda
import auditoria
import classificador
import corpus
import db
import indicadores
import inferencia
import locais as mod_locais
import prompts
import recebidas
import servico
import treinamento
import treino_em_curso
import vocabulario
from classificador import ErroClassificacao

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "chave-de-desenvolvimento")

# Garante o esquema e as colunas acrescentadas depois, aqui e não só no
# `__main__`: quem sobe a aplicação por `flask run` não passaria pelo bloco de
# baixo e encontraria a base no formato antigo.
db.iniciar()

# A lista de locais NÃO é mais carregada em memória na importação. Com 5.571
# municípios, mandá-la inteira para o template seria ~250 KB de HTML em toda
# página que tem o formulário; o navegador pede por estado, sob demanda.
# O estado não é digitado: sai da cidade escolhida, porque uma cidade só
# pertence a um estado e deixar os dois campos soltos convida à contradição.


ROTULOS_STATUS = {
    "rascunho": "Em edição",
    "nao_aprovado": "Não aprovado",
    "aprovado": "Aprovado",
    "unificada": "Unificada",
}


def conexao():
    if "con" not in g:
        g.con = db.conectar()
    return g.con


@app.teardown_appcontext
def fechar_conexao(_erro):
    con = g.pop("con", None)
    if con is not None:
        con.close()


@app.context_processor
def variaveis_globais():
    """As listas fechadas do dicionário vão para todos os templates: é o que
    permite que um select da tela e a validação do servidor venham da mesma
    fonte, sem repetir opção em HTML."""
    pendentes = 0
    try:
        pendentes = recebidas.contar(conexao(), {"estado": "na_fila"})
    except Exception:
        # A barra de navegação não pode derrubar a página por causa de um
        # contador: base sem a tabela ainda, ou fora de contexto de requisição.
        pass
    return {
        "recebidas_pendentes": pendentes,
        # A ajuda da tela aberta, resolvida pelo nome da rota. Vai no contexto
        # global porque o botão mora no `base.html`: assim nenhuma tela
        # precisa declarar nada, e uma tela nova sem ajuda é apontada pela
        # auditoria em vez de simplesmente não ter botão.
        "ajuda_tela": ajuda.para(request.endpoint),
        "rotulos_fenomeno": indicadores.rotulos(),
        "rotulos_status": ROTULOS_STATUS,
        "analista": session.get("analista", ""),
        "vocab": {
            "status_vitima": vocabulario.STATUS_VITIMA,
            "generos": vocabulario.GENEROS,
            "cores_raca": vocabulario.CORES_RACA,
            "grupos_etarios": vocabulario.GRUPOS_ETARIOS,
            "tipos_vitima": vocabulario.TIPOS_VITIMA,
            "circunstancias": vocabulario.CIRCUNSTANCIAS,
            # As que uma pessoa pode marcar à mão. `chacina` fica fora: ela é
            # calculada da lista de vítimas, e uma caixa que o sistema desmarca
            # no próximo recálculo só ensina o analista a não confiar na tela.
            "circunstancias_editaveis": [
                (valor, rotulo) for valor, rotulo in vocabulario.CIRCUNSTANCIAS
                if valor not in vocabulario.CIRCUNSTANCIAS_DERIVADAS
            ],
            "circunstancias_derivadas": vocabulario.CIRCUNSTANCIAS_DERIVADAS,
            "locais_vitima": vocabulario.LOCAIS_VITIMA,
            "corporacoes": vocabulario.CORPORACOES,
            "cargos_agente": vocabulario.CARGOS_AGENTE,
            "status_agente": vocabulario.STATUS_AGENTE,
            "cargos_politico": vocabulario.CARGOS_POLITICO,
            "status_politico": vocabulario.STATUS_POLITICO,
            "faixas_horarias": vocabulario.FAIXAS_HORARIAS,
            "estados": vocabulario.ESTADOS,
            "contextos_acao": vocabulario.CONTEXTOS_ACAO_POLICIAL,
            "contextos_crime": vocabulario.CONTEXTOS_CRIME_ORGANIZADO,
        },
        "rot": {
            "status_vitima": vocabulario.ROTULOS_STATUS_VITIMA,
            "genero": vocabulario.ROTULOS_GENERO,
            "cor_raca": vocabulario.ROTULOS_COR_RACA,
            "grupo_etario": vocabulario.ROTULOS_GRUPO_ETARIO,
            "tipo_vitima": vocabulario.ROTULOS_TIPO_VITIMA,
            "circunstancia": vocabulario.ROTULOS_CIRCUNSTANCIA,
            "local_vitima": vocabulario.ROTULOS_LOCAL_VITIMA,
            "corporacao": vocabulario.ROTULOS_CORPORACAO,
            "cargo_agente": vocabulario.ROTULOS_CARGO_AGENTE,
            "status_agente": vocabulario.ROTULOS_STATUS_AGENTE,
            "cargo_politico": vocabulario.ROTULOS_CARGO_POLITICO,
            "status_politico": vocabulario.ROTULOS_STATUS_POLITICO,
            "faixa_horaria": vocabulario.ROTULOS_FAIXA_HORARIA,
            "estado": vocabulario.ROTULOS_ESTADO,
            "contexto_acao": vocabulario.ROTULOS_CONTEXTO_ACAO,
            "contexto_crime": vocabulario.ROTULOS_CONTEXTO_CRIME,
        },
    }


def analista_atual() -> str:
    return session.get("analista") or "não identificado"


def buscar_ocorrencia(ocorrencia_id: int):
    linha = conexao().execute(
        "SELECT * FROM ocorrencia WHERE id = ?", (ocorrencia_id,)
    ).fetchone()
    if linha is None:
        return None
    ocorrencia = {k: linha[k] for k in linha.keys()}
    ocorrencia["fenomenos"] = db.json_carregar(linha["fenomenos"], [])
    ocorrencia["fenomenos_bert"] = db.json_carregar(linha["fenomenos_bert"], None)
    ocorrencia["nao_sustentados"] = db.json_carregar(linha["nao_sustentados"], [])
    ocorrencia["sugeridos_pela_llm"] = db.json_carregar(
        linha["sugeridos_pela_llm"], []
    )
    ocorrencia["divergencias"] = db.json_carregar(linha["divergencias"], [])
    return ocorrencia


POR_PAGINA = 20


@app.route("/")
def inicio():
    con = conexao()
    status = request.args.get("status", "")
    de = request.args.get("de", "")
    ate = request.args.get("ate", "")

    filtros = ["o.unificada_em IS NULL"]
    parametros = []
    if status:
        filtros.append("o.status = ?")
        parametros.append(status)
    if de:
        filtros.append("o.data_caso >= ?")
        parametros.append(de)
    if ate:
        filtros.append("o.data_caso <= ?")
        parametros.append(ate)
    onde = " AND ".join(filtros)

    total = con.execute(
        f"SELECT COUNT(*) AS total FROM ocorrencia o WHERE {onde}", parametros
    ).fetchone()["total"]

    paginas = max(1, -(-total // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1

    linhas = con.execute(
        "SELECT o.*, (SELECT COUNT(*) FROM vitima v WHERE v.ocorrencia_id = o.id)"
        " AS total_vitimas, (SELECT COUNT(*) FROM noticia n WHERE n.ocorrencia_id = o.id)"
        f" AS total_noticias FROM ocorrencia o WHERE {onde}"
        " ORDER BY o.data_caso DESC, o.id DESC LIMIT ? OFFSET ?",
        [*parametros, POR_PAGINA, (pagina - 1) * POR_PAGINA],
    ).fetchall()

    ocorrencias = []
    for linha in linhas:
        registro = {k: linha[k] for k in linha.keys()}
        registro["fenomenos"] = db.json_carregar(linha["fenomenos"], [])
        ocorrencias.append(registro)

    alertas_abertos = con.execute(
        "SELECT COUNT(*) AS total FROM alerta_duplicidade WHERE status = 'aberto'"
    ).fetchone()["total"]

    return render_template(
        "inicio.html",
        ocorrencias=ocorrencias,
        status_filtro=status,
        de=de,
        ate=ate,
        alertas_abertos=alertas_abertos,
        total=total,
        pagina=pagina,
        paginas=paginas,
    )


TABELAS_EXPORTADAS = [
    "ocorrencia", "noticia", "acao_policial", "crime_organizado",
    "vitimizacao_policial", "politico_baleado", "vitima_historico",
    "alerta_duplicidade",
]

# A planilha de vítimas carrega a localização da ocorrência para que dê para
# buscar por cidade, bairro e localidade sem precisar cruzar as abas na mão.
# Cada coluna diz também como se traduz para texto legível: a planilha sai com
# o rótulo do dicionário, não com o slug do banco, porque quem a abre é uma
# pessoa e não o código.
COLUNAS_PLANILHA_VITIMA = [
    ("id_vitima", "ID_VÍTIMA", None),
    ("id_caso", "ID_CASO", None),
    ("data_caso", "Data do caso", None),
    ("hora_caso", "Hora do caso", None),
    ("estado", "Estado", "estado"),
    ("cidade", "Cidade", None),
    ("bairro", "Bairro", None),
    ("localidade", "Localidade", None),
    ("nome", "Nome da vítima", None),
    ("status", "Status", "status_vitima"),
    ("genero", "Gênero", "genero"),
    ("cor_raca", "Cor/raça", "cor_raca"),
    ("grupo_etario", "Grupo etário", "grupo_etario"),
    ("idade", "Idade", None),
    ("tipo_vitima", "Tipo de vítima", "tipo_vitima"),
    ("circunstancias", "Circunstância", "circunstancia"),
    ("local", "Local", "local_vitima"),
    ("descricao", "Descrição da vítima", None),
    ("data_morte", "Data da morte", None),
    ("fonte", "Fonte", None),
    ("status_ocorrencia", "Status do registro", None),
]

# Colunas das complementares, achatadas na mesma linha da vítima: numa
# planilha, cruzar abas na mão é o que ninguém faz.
COLUNAS_COMPLEMENTO_VITIMA = [
    ("vitimizacao_policial", "corporacao", "Corporação do agente", "corporacao"),
    ("vitimizacao_policial", "cargo", "Cargo do agente", "cargo_agente"),
    ("vitimizacao_policial", "status_agente", "Status do agente", "status_agente"),
    ("vitimizacao_policial", "unidade", "Unidade do agente", None),
    ("politico_baleado", "partido", "Partido político", None),
    ("politico_baleado", "cargo", "Cargo político", "cargo_politico"),
    ("politico_baleado", "status_politico", "Status do político", "status_politico"),
]

_ROTULOS = {
    "estado": vocabulario.ROTULOS_ESTADO,
    "status_vitima": vocabulario.ROTULOS_STATUS_VITIMA,
    "genero": vocabulario.ROTULOS_GENERO,
    "cor_raca": vocabulario.ROTULOS_COR_RACA,
    "grupo_etario": vocabulario.ROTULOS_GRUPO_ETARIO,
    "tipo_vitima": vocabulario.ROTULOS_TIPO_VITIMA,
    "circunstancia": vocabulario.ROTULOS_CIRCUNSTANCIA,
    "local_vitima": vocabulario.ROTULOS_LOCAL_VITIMA,
    "corporacao": vocabulario.ROTULOS_CORPORACAO,
    "cargo_agente": vocabulario.ROTULOS_CARGO_AGENTE,
    "status_agente": vocabulario.ROTULOS_STATUS_AGENTE,
    "cargo_politico": vocabulario.ROTULOS_CARGO_POLITICO,
    "status_politico": vocabulario.ROTULOS_STATUS_POLITICO,
}


def _legivel(valor, dicionario: str | None):
    """Traduz slug para rótulo, e lista para texto separado por vírgula.

    `nao_identificado` é tratado à parte porque também aparece em campo de
    texto livre (bairro, localidade), onde não há lista fechada que o traduza.
    """
    if valor is None:
        return None
    if valor == "nao_identificado":
        return "Não identificado"
    rotulos = _ROTULOS.get(dicionario, {}) if dicionario else {}
    if isinstance(valor, list):
        return ", ".join(rotulos.get(v, v) for v in valor)
    return rotulos.get(valor, valor) if rotulos else valor


def _aba_de_vitimas(planilha, vitimas: list[dict]) -> None:
    aba = planilha.create_sheet("vitima")
    aba.append(
        [rotulo for _, rotulo, _ in COLUNAS_PLANILHA_VITIMA]
        + [rotulo for _, _, rotulo, _ in COLUNAS_COMPLEMENTO_VITIMA]
    )
    for vitima in vitimas:
        linha = [
            _legivel(vitima.get(campo), dicionario)
            for campo, _, dicionario in COLUNAS_PLANILHA_VITIMA
        ]
        for chave, coluna, _, dicionario in COLUNAS_COMPLEMENTO_VITIMA:
            linha.append(
                _legivel((vitima.get(chave) or {}).get(coluna), dicionario)
            )
        aba.append(linha)


def _enviar_planilha(planilha, nome: str):
    buffer = io.BytesIO()
    planilha.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"{nome}-{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/exportar")
def exportar():
    """Baixa a base inteira como uma planilha, uma aba por tabela."""
    from openpyxl import Workbook

    con = conexao()
    planilha = Workbook()
    planilha.remove(planilha.active)

    for tabela in TABELAS_EXPORTADAS:
        aba = planilha.create_sheet(tabela)
        colunas = [c[1] for c in con.execute(f"PRAGMA table_info({tabela})")]
        aba.append(colunas)
        for linha in con.execute(f"SELECT * FROM {tabela}"):
            aba.append([linha[coluna] for coluna in colunas])

    _aba_de_vitimas(planilha, servico.filtrar_vitimas(con, {}))
    return _enviar_planilha(planilha, "base-classificador")


FILTROS_VITIMA = (
    "de", "ate", "estado", "cidade", "bairro", "localidade", "circunstancia",
    "status", "genero", "cor_raca", "grupo_etario", "tipo_vitima",
)


def _filtros_de_vitima() -> dict:
    filtros = {
        chave: request.args.get(chave, "").strip() for chave in FILTROS_VITIMA
    }
    # O acompanhamento de feridos é um preset desta mesma tela, não outra
    # tela: a tabela de vítimas é uma só, e duplicá-la faria a cor/raça ser
    # editável num lugar e não no outro.
    if request.args.get("acompanhamento"):
        filtros = servico.filtros_de_acompanhamento(filtros)
    return filtros


@app.route("/vitimas")
def vitimas():
    con = conexao()
    filtros = _filtros_de_vitima()
    total = servico.contar_vitimas(con, filtros)

    paginas = max(1, -(-total // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1

    return render_template(
        "vitimas.html",
        vitimas=servico.filtrar_vitimas(con, filtros, POR_PAGINA, (pagina - 1) * POR_PAGINA),
        filtros=filtros,
        # Sem os vazios, para não sujar as URLs de paginação e download.
        filtros_ativos={chave: valor for chave, valor in filtros.items() if valor},
        total=total,
        pagina=pagina,
        paginas=paginas,
        # As opções do filtro saem do que EXISTE no banco, não do catálogo
        # de 5.571 municípios: filtrar por uma cidade sem nenhuma vítima
        # devolveria sempre vazio, e procurá-la numa lista de milhares para
        # descobrir isso é trabalho jogado fora.
        locais=_locais_com_registro(con),
        em_acompanhamento=bool(request.args.get("acompanhamento")),
        dias_acompanhamento=servico.DIAS_ACOMPANHAMENTO,
        partidos=vocabulario.partidos(),
    )


def _locais_com_registro(con) -> dict[str, list[str]]:
    """Cidade -> bairros, apenas onde já há vítima registrada.

    Alimenta os filtros das telas de consulta. Aqui a repetição de nome de
    município entre estados não atrapalha: quem filtra está procurando no que
    foi registrado, e o registro guarda a UF junto.
    """
    mapa: dict[str, list[str]] = {}
    for linha in con.execute(
        "SELECT DISTINCT o.cidade, o.bairro FROM ocorrencia o"
        " JOIN vitima v ON v.ocorrencia_id = o.id"
        " WHERE o.cidade IS NOT NULL AND o.cidade != ''"
        " ORDER BY o.cidade, o.bairro"
    ):
        bairros = mapa.setdefault(linha["cidade"], [])
        if linha["bairro"] and linha["bairro"] not in bairros:
            bairros.append(linha["bairro"])
    return mapa


@app.route("/vitima/<int:vitima_id>/rapida", methods=["POST"])
def editar_vitima_rapida(vitima_id: int):
    """Edição direto na tabela de Vítimas. Serve ao que não vem da notícia:
    cor/raça quase nunca é informada pela imprensa, e o acompanhamento de 90
    dias existe justamente para corrigir o status de quem morreu depois."""
    con = conexao()
    dados = {
        campo: request.form.get(campo)
        for campo in ("cor_raca", "status", "data_morte", "descricao")
        if campo in request.form
    }
    with con:
        servico.atualizar_vitima(
            con, vitima_id, dados, analista_atual(),
            request.form.get("fonte_alteracao", "") or "Edição na tela de Vítimas",
        )
    flash("Vítima atualizada.", "ok")
    return redirect(request.referrer or url_for("vitimas"))


@app.route("/vitimas/exportar")
def exportar_vitimas():
    """Planilha só de vítimas, respeitando a busca ativa na tela."""
    from openpyxl import Workbook

    planilha = Workbook()
    planilha.remove(planilha.active)
    _aba_de_vitimas(planilha, servico.filtrar_vitimas(conexao(), _filtros_de_vitima()))
    return _enviar_planilha(planilha, "vitimas")


@app.route("/analista", methods=["POST"])
def definir_analista():
    session["analista"] = request.form.get("analista", "").strip()
    return redirect(request.referrer or url_for("inicio"))


@app.route("/ocorrencia/nova", methods=["GET", "POST"])
def nova_ocorrencia():
    if request.method == "POST":
        nome = request.form.get("analista", "").strip()
        if nome:
            session["analista"] = nome
        con = conexao()
        momento = db.agora()
        hora = request.form.get("hora_caso", "").strip() or None
        cidade = request.form["cidade"]
        estado = request.form.get("estado", "").strip().upper()
        # O par estado+cidade é conferido no servidor: um `select` adulterado
        # no navegador gravaria uma cidade que não existe naquela UF, e o par
        # errado contamina toda estatística por local.
        if not mod_locais.existe_cidade(estado, cidade):
            flash(
                f"{cidade} não consta em {estado or 'nenhum estado'}."
                " Escolha pela lista, ou acrescente a cidade em Locais.",
                "erro",
            )
            return redirect(url_for("nova_ocorrencia"))
        with con:
            cursor = con.execute(
                "INSERT INTO ocorrencia (data_registro, data_caso, hora_caso,"
                " faixa_horaria, estado, cidade, bairro, localidade, analista,"
                " atualizado_em, status)"
                " VALUES (?,?,?,?,?,?,?,?,?,?,'rascunho')",
                (
                    momento,
                    request.form["data_caso"],
                    hora,
                    # Derivada da hora, nunca digitada: dois campos que podem
                    # se contradizer sempre acabam se contradizendo.
                    vocabulario.faixa_horaria(hora),
                    estado,
                    cidade,
                    request.form.get("bairro", "").strip() or "nao_identificado",
                    request.form.get("localidade", "").strip() or "nao_identificado",
                    analista_atual(),
                    momento,
                ),
            )
            ocorrencia_id = cursor.lastrowid

            # Veio da fila: o texto já está na mão, então a notícia é anexada
            # aqui mesmo. Sem isto o analista teria de copiar e colar de novo o
            # que acabou de ler, e a recebida ficaria sem saber que virou algo.
            vinda_da_fila = request.form.get("recebida_id", "").strip()
            if vinda_da_fila.isdigit():
                registro = recebidas.buscar(con, int(vinda_da_fila))
                if registro and registro["estado"] != "virou_ocorrencia":
                    servico.adicionar_noticia(
                        con, ocorrencia_id, registro["link"] or "",
                        registro["texto"],
                    )
                    recebidas.marcar_como_ocorrencia(
                        con, registro["id"], ocorrencia_id
                    )

        if vinda_da_fila.isdigit():
            # A classificação fica para o botão de reprocessar, como em qualquer
            # ocorrência: aqui ela custaria uma chamada à LLM antes de o
            # analista conferir data e endereço, que é o que ele acabou de
            # digitar e ainda não viu gravado.
            flash(
                "Ocorrência criada a partir da fila, com a notícia já anexada."
                " Reprocesse para classificar.",
                "ok",
            )
        return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))

    # Vindo da fila, a data e o texto já são conhecidos: o que falta é o
    # endereço, que é o que esta tela existe para pedir.
    da_fila = None
    pedido = request.args.get("recebida", "").strip()
    if pedido.isdigit():
        da_fila = recebidas.buscar(conexao(), int(pedido))

    return render_template(
        "nova_ocorrencia.html", estados=mod_locais.estados(),
        hoje=date.today(), da_fila=da_fila,
    )


@app.route("/ocorrencia/<int:ocorrencia_id>")
def ocorrencia(ocorrencia_id: int):
    con = conexao()
    registro = buscar_ocorrencia(ocorrencia_id)
    if registro is None:
        flash("Ocorrência não encontrada.", "erro")
        return redirect(url_for("inicio"))

    noticias = con.execute(
        "SELECT * FROM noticia WHERE ocorrencia_id = ? ORDER BY id", (ocorrencia_id,)
    ).fetchall()
    vitimas = servico.listar_vitimas(con, ocorrencia_id)
    historico = con.execute(
        "SELECT * FROM vitima_historico WHERE ocorrencia_id = ?"
        " ORDER BY alterado_em DESC, id DESC",
        (ocorrencia_id,),
    ).fetchall()
    alertas = con.execute(
        "SELECT * FROM alerta_duplicidade WHERE status = 'aberto'"
        " AND (ocorrencia_id = ? OR ocorrencia_similar_id = ?)",
        (ocorrencia_id, ocorrencia_id),
    ).fetchall()

    return render_template(
        "ocorrencia.html",
        ocorrencia=registro,
        noticias=noticias,
        vitimas=vitimas,
        historico=historico,
        alertas=alertas,
        complementos=servico.complementos_da_ocorrencia(con, ocorrencia_id),
        origem_fenomenos=indicadores.origem(vitimas, registro["fenomenos"]),
        partidos=vocabulario.partidos(),
    )


@app.route("/ocorrencia/<int:ocorrencia_id>/noticia", methods=["POST"])
def adicionar_noticia(ocorrencia_id: int):
    texto = request.form.get("texto", "").strip()
    if not texto:
        flash("O texto da notícia é obrigatório.", "erro")
        return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))

    con = conexao()
    # Antes de gravar: esta mesma notícia já está em outro registro? Não
    # bloqueia — às vezes duas ocorrências próximas citam a mesma reportagem —
    # mas o analista precisa ver, porque anexar a notícia errada faz as vítimas
    # de um fato entrarem na lista de outro.
    repetida = servico.noticia_em_outra_ocorrencia(con, texto, ocorrencia_id)

    with con:
        servico.adicionar_noticia(
            con, ocorrencia_id, request.form.get("link", "").strip(), texto
        )

    if repetida:
        onde = ", ".join(
            f"{r['id_caso'] or r['ocorrencia_id']}" for r in repetida
        )
        flash(
            f"Atenção: este mesmo texto já está anexado em {onde}. Se forem"
            " fatos diferentes, remova a notícia daqui — senão as vítimas de um"
            " vão parar na lista do outro.",
            "aviso",
        )

    try:
        with con:
            servico.reprocessar(con, ocorrencia_id)
            servico.marcar_para_reaprovacao(con, ocorrencia_id)
        flash(
            "Notícia adicionada. Todas as notícias foram relidas juntas e a lista"
            " de vítimas foi refeita — confira antes de concluir.",
            "ok",
        )
    except ErroClassificacao as erro:
        flash(f"A notícia foi salva, mas a classificação falhou: {erro}", "erro")

    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/ocorrencia/<int:ocorrencia_id>/reprocessar", methods=["POST"])
def reprocessar(ocorrencia_id: int):
    con = conexao()
    try:
        with con:
            servico.reprocessar(con, ocorrencia_id)
        flash("Ocorrência reprocessada.", "ok")
    except ErroClassificacao as erro:
        flash(str(erro), "erro")
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/noticia/<int:noticia_id>/excluir", methods=["POST"])
def excluir_noticia(noticia_id: int):
    con = conexao()
    linha = con.execute(
        "SELECT ocorrencia_id FROM noticia WHERE id = ?", (noticia_id,)
    ).fetchone()
    # Notícia já excluída — por duplo envio do formulário ou por link velho.
    # Sem isto, o `linha["ocorrencia_id"]` abaixo estoura em 500.
    if linha is None:
        abort(404)
    ocorrencia_id = linha["ocorrencia_id"]
    with con:
        con.execute("DELETE FROM noticia WHERE id = ?", (noticia_id,))
    restantes = con.execute(
        "SELECT COUNT(*) AS total FROM noticia WHERE ocorrencia_id = ?", (ocorrencia_id,)
    ).fetchone()["total"]
    if restantes:
        try:
            with con:
                servico.reprocessar(con, ocorrencia_id)
        except ErroClassificacao as erro:
            flash(str(erro), "erro")
    else:
        with con:
            servico.limpar_classificacao(con, ocorrencia_id)
        flash("Última notícia removida: a classificação e as vítimas foram limpas.", "aviso")
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


def _dados_de_vitima() -> dict:
    """Lê o formulário de vítima. Campo de marcação múltipla vem por getlist;
    os das tabelas complementares vêm prefixados pela tabela."""
    dados = {campo: request.form.get(campo) for campo in servico.CAMPOS_VITIMA}
    for campo in servico.CAMPOS_VITIMA_LISTA:
        dados[campo] = request.form.getlist(campo)
    for chave, config in servico.COMPLEMENTOS_VITIMA.items():
        for coluna in config["colunas"]:
            nome = f"{chave}_{coluna}"
            if nome in request.form:
                dados[nome] = request.form.get(nome)
    return dados


@app.route("/vitima/<int:vitima_id>", methods=["POST"])
def editar_vitima(vitima_id: int):
    con = conexao()
    with con:
        ocorrencia_id = servico.atualizar_vitima(
            con, vitima_id, _dados_de_vitima(), analista_atual(),
            request.form.get("fonte_alteracao", ""),
        )
    flash("Vítima atualizada. Os fenômenos foram recalculados.", "ok")
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/vitima/<int:vitima_id>/proposta/<acao>", methods=["POST"])
def resolver_proposta(vitima_id: int, acao: str):
    con = conexao()
    with con:
        if acao == "aceitar":
            ocorrencia_id = servico.aceitar_proposta(con, vitima_id, analista_atual())
        else:
            ocorrencia_id = servico.descartar_proposta(con, vitima_id)
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/ocorrencia/<int:ocorrencia_id>/vitima", methods=["POST"])
def criar_vitima(ocorrencia_id: int):
    con = conexao()
    with con:
        servico.criar_vitima_manual(
            con, ocorrencia_id, _dados_de_vitima(), analista_atual()
        )
    flash("Vítima incluída manualmente.", "ok")
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/vitima/<int:vitima_id>/excluir", methods=["POST"])
def excluir_vitima(vitima_id: int):
    con = conexao()
    with con:
        ocorrencia_id = servico.remover_vitima(con, vitima_id, analista_atual())
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/ocorrencia/<int:ocorrencia_id>/concluir", methods=["POST"])
def concluir(ocorrencia_id: int):
    con = conexao()
    total = con.execute(
        "SELECT COUNT(*) AS total FROM noticia WHERE ocorrencia_id = ?", (ocorrencia_id,)
    ).fetchone()["total"]
    if not total:
        flash("Inclua ao menos uma notícia antes de concluir.", "erro")
        return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))

    with con:
        servico.recalcular_fenomenos(con, ocorrencia_id)
        con.execute(
            "UPDATE ocorrencia SET status = 'nao_aprovado', atualizado_em = ?"
            " WHERE id = ?",
            (db.agora(), ocorrencia_id),
        )
        similares = servico.checar_duplicidade(con, ocorrencia_id)

    if similares:
        flash(
            f"Registro gravado. Atenção: {len(similares)} possível(is) duplicidade(s)"
            " no mesmo dia e bairro — verifique em Alertas.",
            "aviso",
        )
    else:
        flash("Registro gravado como não aprovado e enviado para revisão.", "ok")
    return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))


@app.route("/aprovacao")
def aprovacao():
    con = conexao()
    linhas = con.execute(
        "SELECT o.*, (SELECT COUNT(*) FROM alerta_duplicidade a WHERE a.status = 'aberto'"
        " AND (a.ocorrencia_id = o.id OR a.ocorrencia_similar_id = o.id)) AS alertas"
        " FROM ocorrencia o WHERE o.status = 'nao_aprovado' AND o.unificada_em IS NULL"
        " ORDER BY o.atualizado_em"
    ).fetchall()
    registros = []
    for linha in linhas:
        registro = {k: linha[k] for k in linha.keys()}
        registro["fenomenos"] = db.json_carregar(linha["fenomenos"], [])
        registros.append(registro)
    return render_template("aprovacao.html", ocorrencias=registros)


@app.route("/ocorrencia/<int:ocorrencia_id>/aprovar", methods=["POST"])
def aprovar(ocorrencia_id: int):
    con = conexao()
    aberto = con.execute(
        "SELECT COUNT(*) AS total FROM alerta_duplicidade WHERE status = 'aberto'"
        " AND (ocorrencia_id = ? OR ocorrencia_similar_id = ?)",
        (ocorrencia_id, ocorrencia_id),
    ).fetchone()["total"]
    if aberto:
        flash(
            "Existe alerta de duplicidade aberto para este registro. Resolva o"
            " alerta antes de aprovar.",
            "erro",
        )
        return redirect(url_for("ocorrencia", ocorrencia_id=ocorrencia_id))

    with con:
        con.execute(
            "UPDATE ocorrencia SET status = 'aprovado', aprovado_por = ?,"
            " aprovado_em = ?, observacao_aprovacao = ?, atualizado_em = ? WHERE id = ?",
            (
                analista_atual(), db.agora(),
                request.form.get("observacao", "").strip() or None,
                db.agora(), ocorrencia_id,
            ),
        )
        # O ciclo fecha aqui: a conferência que o analista acabou de fazer vira
        # o rótulo do próximo treino. Vai na MESMA transação da aprovação, para
        # que não exista registro aprovado sem o exemplo correspondente.
        #
        # A caixa "servir de exemplo de treino" vem MARCADA no formulário, e
        # desmarcá-la grava a notícia no corpus sem aprovação, em vez de não
        # gravar: o registro está certo, só não serve de exemplo — texto
        # truncado, duas ocorrências misturadas, fonte duvidosa. Ela espera
        # revisão em /corpus em lugar de treinar o modelo ou de se perder.
        treinar = request.form.get("treinar", "1") in ("1", "on", "sim")
        enviadas = servico.enviar_ao_corpus(
            con, ocorrencia_id, analista_atual(), aprovar=treinar
        )

    recado = "Registro aprovado e liberado para a comunicação."
    if enviadas["novas"]:
        rotulos = ", ".join(enviadas["rotulos"]) or "nenhum indicador"
        destino = (
            "foram ao corpus de treino" if treinar
            else "foram ao corpus SEM aprovação, esperando revisão em /corpus"
        )
        recado += (
            f" {len(enviadas['novas'])} notícia(s) {destino}"
            f" com os rótulos: {rotulos}."
        )
        if enviadas["repetidas"]:
            recado += f" {len(enviadas['repetidas'])} já estavam lá."
    elif enviadas["repetidas"]:
        recado += " As notícias já estavam no corpus de treino."
    flash(recado, "ok")
    return redirect(url_for("aprovacao"))


@app.route("/corpus/<int:noticia_id>/aprovar", methods=["POST"])
def corpus_aprovar(noticia_id: int):
    """O sênior assume o rótulo de um exemplo, ou o devolve para revisão.

    Separado da edição de propósito: assumir um rótulo que já está certo é um
    clique, não um formulário inteiro. É esta tela que esvazia a fila de
    exemplos que o júnior rotulou.
    """
    aprovar_agora = request.form.get("aprovada", "1") in ("1", "on", "sim")
    con = conexao()
    with con:
        corpus.aprovar(con, noticia_id, analista_atual(), aprovada=aprovar_agora)
    flash(
        f"Notícia #{noticia_id} "
        + ("aprovada: passa a treinar o modelo."
           if aprovar_agora else
           "devolvida para revisão: sai do treino até ser aprovada."),
        "ok",
    )
    return redirect(request.referrer or url_for("corpus_repositorio"))


@app.route("/ocorrencia/<int:ocorrencia_id>/reprovar", methods=["POST"])
def reprovar(ocorrencia_id: int):
    con = conexao()
    with con:
        con.execute(
            "UPDATE ocorrencia SET status = 'nao_aprovado', aprovado_por = NULL,"
            " aprovado_em = NULL, observacao_aprovacao = ?, atualizado_em = ?"
            " WHERE id = ?",
            (
                request.form.get("observacao", "").strip() or None,
                db.agora(), ocorrencia_id,
            ),
        )
    flash("Registro devolvido ao analista para ajuste.", "aviso")
    return redirect(url_for("aprovacao"))


@app.route("/alertas")
def alertas():
    con = conexao()
    linhas = con.execute(
        "SELECT a.*, o1.data_caso, o1.cidade, o1.bairro,"
        " o1.analista AS analista_a, o2.analista AS analista_b"
        " FROM alerta_duplicidade a"
        " JOIN ocorrencia o1 ON o1.id = a.ocorrencia_id"
        " JOIN ocorrencia o2 ON o2.id = a.ocorrencia_similar_id"
        " ORDER BY a.status = 'aberto' DESC, a.criado_em DESC"
    ).fetchall()
    return render_template("alertas.html", alertas=linhas)


@app.route("/alerta/<int:alerta_id>/<acao>", methods=["POST"])
def resolver_alerta(alerta_id: int, acao: str):
    con = conexao()
    motivo = request.form.get("motivo", "").strip()
    try:
        with con:
            if acao == "unificar":
                destino = servico.unificar(con, alerta_id, analista_atual(), motivo)
                flash(
                    f"Registros unificados na ocorrência #{destino}. As notícias"
                    " passaram de novo pela LLM e o registro voltou para aprovação.",
                    "ok",
                )
            elif acao == "desfazer":
                servico.desfazer_unificacao(con, alerta_id, analista_atual())
                flash("Unificação desfeita. Os registros voltaram a ser separados.", "ok")
            else:
                servico.descartar_alerta(con, alerta_id, analista_atual(), motivo)
                flash("Alerta descartado: os registros seguem separados.", "ok")
    except ErroClassificacao as erro:
        flash(str(erro), "erro")
    return redirect(url_for("alertas"))


@app.route("/acompanhamento")
def acompanhamento():
    """Mantida como atalho: o acompanhamento de feridos passou a ser um preset
    da tela de Vítimas, para que exista uma tabela de vítimas só."""
    return redirect(
        url_for(
            "vitimas",
            acompanhamento=1,
            de=request.args.get("de", "").strip() or None,
            ate=request.args.get("ate", "").strip() or None,
        )
    )


@app.route("/gestao", methods=["GET", "POST"])
def gestao():
    if request.method == "POST":
        lista = indicadores.carregar()
        nome = request.form.get("nome", "").strip().lower().replace(" ", "_")
        if not nome:
            flash("Informe o nome (slug) do indicador.", "erro")
            return redirect(url_for("gestao"))

        # Parte do indicador que já existe e escreve o formulário SOBRE ele. O
        # contrário — montar um dicionário novo — apagava calado os campos que a
        # tela não enviava: `decidido_por` (e o indicador saía da lista de
        # classes do BERT), `classe`, `circunstancias` (e caía o portão do
        # feminicídio) e `saidas` (e caíam as tabelas complementares).
        atual = next((i for i in lista if i["nome"] == nome), None)
        indicador = dict(atual) if atual else {}

        # O formulário pré-preenchido pelo botão "Editar" manda o estado
        # completo: ali, o que volta vazio foi desmarcado de propósito. No
        # formulário em branco, vazio significa "não mexi nisso", e o valor
        # antigo é mantido — é o que impede um cadastro apressado de zerar a
        # configuração de um indicador que já existia.
        editando = request.form.get("editando") == "1"

        indicador["nome"] = nome
        indicador["rotulo"] = request.form.get("rotulo", "").strip() or nome
        # Todo indicador é classificado a partir do texto: é essa a lista
        # de classes do modelo. O tipo continua no arquivo por compatibilidade.
        indicador["tipo"] = "texto"
        indicador["definicao"] = request.form.get("definicao", "").strip()
        indicador["skill"] = request.form.get("skill", "").strip()

        classe = request.form.get("classe", "").strip()
        if classe in ("fenomeno", "circunstancia"):
            indicador["classe"] = classe
        decisor = request.form.get("decidido_por", "").strip()
        if decisor in ("bert", "llm", "regra"):
            indicador["decidido_por"] = decisor

        for campo, enviados, validos in (
            ("circunstancias", request.form.getlist("circunstancias"),
             [valor for valor, _ in vocabulario.CIRCUNSTANCIAS]),
        ):
            escolhidos = [slug for slug in enviados if slug in validos]
            anteriores = indicador.get(campo)
            if not escolhidos and (not editando or anteriores is None):
                # Vazio e sem edição explícita: não mexe. Vazio numa chave que
                # nunca existiu: não a cria — salvar sem mexer tem de deixar o
                # arquivo idêntico.
                continue
            # Mantém a ordem que o arquivo já tinha para os que continuam
            # marcados e põe os novos no fim, em vez de reordenar tudo pela
            # ordem dos checkboxes da tela.
            mantidos = [slug for slug in (anteriores or []) if slug in set(escolhidos)]
            indicador[campo] = mantidos + [
                slug for slug in escolhidos if slug not in set(anteriores or [])
            ]

        exemplos = [
            linha.strip()
            for linha in request.form.get("exemplos", "").splitlines()
            if linha.strip()
        ]
        if exemplos or (editando and indicador.get("exemplos") is not None):
            indicador["exemplos"] = exemplos

        # `saidas` é JSON aninhado (chave, formato, notas, por_vitima). Um
        # formulário campo a campo para isso seria frágil e ainda assim
        # incompleto; o texto cru é editável e verificável. JSON inválido ABORTA
        # o salvamento, para que um erro de digitação não apague as tabelas.
        bruto = request.form.get("saidas", "").strip()
        if bruto:
            try:
                decodificado = json.loads(bruto)
            except json.JSONDecodeError as erro:
                flash(
                    f"As tabelas complementares não são um JSON válido: {erro}."
                    " Nada foi salvo.", "erro",
                )
                return redirect(url_for("gestao", editar=nome))
            if not isinstance(decodificado, list):
                flash(
                    "As tabelas complementares devem ser uma LISTA JSON, mesmo"
                    " com uma tabela só. Nada foi salvo.", "erro",
                )
                return redirect(url_for("gestao", editar=nome))
            indicador["saidas"] = decodificado
        elif editando:
            indicador.pop("saidas", None)

        # A regra é opcional e não substitui a classificação: ela confere o
        # indicador depois, sobre a lista de vítimas que a LLM extraiu.
        if request.form.get("tem_regra"):
            # Parte da regra atual, e não de um dicionário vazio: `campo_em` usa
            # a lista `valores`, que esta tela não edita e que seria perdida.
            regra = dict(indicador.get("regra") or {})
            regra["operacao"] = request.form.get("operacao", "campo_igual")
            regra["campo"] = request.form.get("campo", "").strip()

            valor = request.form.get("valor", "").strip()
            if valor and regra["operacao"] == "campo_menor_que":
                if not valor.lstrip("-").isdigit():
                    flash(
                        f"A operação campo_menor_que precisa de um número, e"
                        f" '{valor}' não é. Nada foi salvo.", "erro",
                    )
                    return redirect(url_for("gestao", editar=nome))
                regra["valor"] = int(valor)
            elif valor:
                regra["valor"] = valor
            elif editando:
                regra.pop("valor", None)

            minimo = request.form.get("minimo", "").strip()
            if minimo.isdigit():
                regra["minimo"] = int(minimo)
            elif editando:
                regra.pop("minimo", None)

            indicador["regra"] = regra
        elif editando:
            indicador.pop("regra", None)

        # Substitui NO LUGAR. Reanexar no fim mudava a ordem da lista, e essa
        # ordem define as posições do vetor multirrótulo: trocá-la depois de um
        # treino deslocaria a leitura de todas as classes.
        if atual is not None:
            lista = [indicador if i["nome"] == nome else i for i in lista]
        else:
            # A DATA DE NASCIMENTO do indicador. É o que permite depois achar
            # as notícias anotadas antes dele existir, que no treino contariam
            # como exemplo negativo de uma pergunta que ninguém lhes fez.
            # Só nos novos: os do catálogo original não têm data e são tratados
            # como se sempre tivessem existido.
            indicador["criado_em"] = db.agora()
            lista.append(indicador)

        indicadores.salvar(lista)
        recado = f"Indicador '{nome}' salvo."
        if atual is None and indicador.get("decidido_por") == "bert":
            antigas = corpus.contar(conexao(), {"rotulada": "sim"})
            recado += (
                " Ele só passa a ser marcado pelo BERT depois de receber"
                f" {corpus.MINIMO_POR_CLASSE} positivos APROVADOS no corpus e de"
                " um novo treino — o modelo em uso foi treinado com a lista"
                " antiga de classes."
            )
            if antigas:
                recado += (
                    f" Atenção: as {antigas} notícia(s) já anotadas contam como"
                    " exemplo NEGATIVO dele até alguém relê-las."
                )
        elif atual is None or indicador.get("decidido_por") == "bert":
            recado += (
                " Ele só passa a ser marcado pelo BERT depois de receber"
                " exemplos rotulados no corpus e de um novo treino."
            )
        flash(
            recado,
            "ok",
        )
        return redirect(url_for("gestao"))

    # Preenche o formulário com um indicador existente, para que editar uma
    # skill não obrigue a redigitar o resto.
    edicao = indicadores.buscar((request.args.get("editar") or "").strip())
    if edicao:
        # Serializado aqui, e não pelo filtro `tojson` do template: o filtro
        # escapa caracteres para uso dentro de <script> e eles apareceriam
        # literais na área de edição.
        edicao = {
            **edicao,
            "saidas_json": json.dumps(
                edicao["saidas"], ensure_ascii=False, indent=2
            ) if edicao.get("saidas") else "",
        }

    return render_template(
        "gestao.html",
        indicadores=indicadores.carregar(),
        circunstancias=vocabulario.CIRCUNSTANCIAS,
        condicionadas=vocabulario.CIRCUNSTANCIAS_CONDICIONADAS,
        edicao=edicao,
    )


@app.route("/prompts", methods=["GET", "POST"])
def editar_prompts():
    """Os prompts que não pertencem a um indicador: o prompt mestre, o formato
    de saída e os três cabeçalhos da tarefa de classes.

    Separado de /gestao de propósito. Ali se edita o prompt de UM indicador, que
    só viaja quando o BERT o marca; aqui se edita o que vai em TODA chamada.
    """
    if request.method == "POST":
        if request.form.get("restaurar"):
            nome = request.form["restaurar"]
            prompts.restaurar(nome)
            flash(
                f"Bloco '{nome}' voltou ao padrão de fábrica.", "ok",
            )
            return redirect(url_for("editar_prompts"))

        textos = {
            nome: request.form.get(nome, "")
            for nome in prompts.PADROES
            if nome in request.form
        }
        vazios = [nome for nome, texto in textos.items() if not texto.strip()]
        if vazios:
            flash(
                "Prompt vazio não é salvo — ele deixaria a LLM sem a instrução:"
                f" {', '.join(vazios)}. Para voltar ao texto original use"
                " 'Restaurar padrão'.",
                "erro",
            )
            return redirect(url_for("editar_prompts"))

        # Apagar um marcador é uma escolha legítima: a LLM passa a não receber
        # aquela lista. Mas é uma escolha que precisa ser consciente, então ela
        # é salva e avisada, não bloqueada.
        avisos = []
        for nome, texto in textos.items():
            faltando = prompts.marcadores_faltando(nome, texto)
            if faltando:
                avisos.append(
                    f"{nome}: sem {', '.join('{' + m + '}' for m in faltando)}"
                )

        prompts.salvar(textos)
        flash("Prompts salvos.", "ok")
        if avisos:
            flash(
                "Atenção, marcador removido — a LLM deixará de receber essa"
                f" lista: {'; '.join(avisos)}.",
                "erro",
            )
        return redirect(url_for("editar_prompts"))

    blocos = prompts.carregar()
    return render_template(
        "prompts.html",
        blocos=blocos,
        ordem=prompts.ORDEM,
        marcadores=prompts.MARCADORES,
        # Os tamanhos de cada situação, para que a economia do filtro seja
        # visível na mesma tela em que os prompts são escritos.
        tamanhos={
            "nucleo": len(classificador.montar_nucleo()),
            "duas_classes": len(
                classificador.montar_prompt_indicadores(
                    indicadores.nomes()[:2]
                )
            ),
            "nenhuma": len(classificador.montar_prompt_indicadores([])),
            "sem_modelo": len(classificador.montar_prompt_indicadores(None)),
        },
    )


@app.route("/gestao/<nome>/excluir", methods=["POST"])
def excluir_indicador(nome: str):
    indicadores.salvar([i for i in indicadores.carregar() if i["nome"] != nome])
    flash(f"Indicador '{nome}' removido.", "ok")
    return redirect(url_for("gestao"))


@app.route("/locais/cidades/<sigla>")
def cidades_do_estado(sigla: str):
    """As cidades de um estado, para o formulário preencher o segundo campo.

    Sob demanda, e não embutido na página: Minas sozinha tem 853 municípios, e
    os 27 estados juntos passariam de 250 KB em cada carregamento."""
    return {"cidades": mod_locais.cidades(sigla)}


@app.route("/locais/bairros/<sigla>/<cidade>")
def bairros_da_cidade(sigla: str, cidade: str):
    return {"bairros": mod_locais.bairros(sigla, cidade)}


# --------------------------------------------------------------------------
# Corpus de treino, treino dos modelos e testes antes da camada de LLM
# --------------------------------------------------------------------------

def _filtros_de_corpus() -> dict:
    return {
        chave: request.args.get(chave, "").strip()
        for chave in ("rotulada", "aprovada", "indicador", "busca", "particao",
                      "reanotar")
    }


@app.route("/corpus")
def corpus_repositorio():
    con = conexao()
    filtros = _filtros_de_corpus()
    total = corpus.contar(con, filtros)

    paginas = max(1, -(-total // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1

    return render_template(
        "corpus.html",
        noticias=corpus.listar(con, filtros, POR_PAGINA, (pagina - 1) * POR_PAGINA),
        filtros=filtros,
        filtros_ativos={c: v for c, v in filtros.items() if v},
        estatisticas=corpus.estatisticas(con),
        catalogo=indicadores.do_bert(),
        # Quantas notícias precisam ser relidas por causa de um indicador
        # criado depois delas. Sem este número, a dívida fica invisível e vira
        # um modelo que nunca marca a classe nova.
        a_reanotar=corpus.a_reanotar(con),
        codigo_divisao=corpus.CODIGO_DIVISAO,
        total=total,
        pagina=pagina,
        paginas=paginas,
    )


@app.route("/corpus/nova", methods=["POST"])
def corpus_adicionar():
    texto = request.form.get("texto", "").strip()
    if not texto:
        flash("O texto da notícia é obrigatório.", "erro")
        return redirect(url_for("corpus_repositorio"))

    con = conexao()
    dados = {
        "link": request.form.get("link"),
        "texto": texto,
        "indicadores": request.form.getlist("indicadores"),
        # No formulario de inclusao nao ha botao de "nenhum se aplica": marcar
        # um indicador ja anota a noticia, nao marcar nenhum manda para a fila.
        "rotulada": request.form.get("rotulada"),
    }
    with con:
        novo = corpus.adicionar(con, dados, analista_atual())
    flash(f"Notícia #{novo} incluída no corpus.", "ok")

    if request.form.get("continuar"):
        return redirect(url_for("corpus_repositorio", _anchor="nova"))
    return redirect(url_for("corpus_noticia", noticia_id=novo))


@app.route("/corpus/<int:noticia_id>", methods=["GET", "POST"])
def corpus_noticia(noticia_id: int):
    con = conexao()

    if request.method == "POST":
        dados = {
            "link": request.form.get("link"),
            "texto": request.form.get("texto", "").strip(),
            "indicadores": request.form.getlist("indicadores"),
            # A tela de anotacao sempre manda rotulada=1: decidir que NENHUM
            # indicador se aplica e uma anotacao, nao uma pendencia.
            "rotulada": "1",
            # A caixa de aprovação vem no mesmo envio, para que corrigir o
            # rótulo e assumir a correção sejam um clique só. Sem ela no form,
            # `atualizar` derrubaria a aprovação a cada correção e o sênior
            # teria de aprovar de novo numa segunda tela.
            "aprovada": "1" if request.form.get("aprovada") else "0",
        }
        if not dados["texto"]:
            flash("O texto da notícia é obrigatório.", "erro")
            return redirect(url_for("corpus_noticia", noticia_id=noticia_id))
        with con:
            corpus.atualizar(con, noticia_id, dados, analista_atual())
        flash(
            "Rótulos salvos e aprovados: esta notícia passa a treinar o modelo."
            if dados["aprovada"] == "1" else
            "Rótulos salvos. A notícia NÃO treina o modelo até ser aprovada.",
            "ok",
        )
        proxima = request.form.get("proxima")
        if proxima:
            return redirect(url_for("corpus_anotar"))
        return redirect(url_for("corpus_noticia", noticia_id=noticia_id))

    registro = corpus.buscar(con, noticia_id)
    if registro is None:
        flash("Notícia não encontrada no corpus.", "erro")
        return redirect(url_for("corpus_repositorio"))

    return render_template(
        "corpus_noticia.html",
        noticia=registro,
        catalogo=indicadores.do_bert(),
        restantes=corpus.contar(con, {"rotulada": "nao"}),
    )


@app.route("/corpus/anotar")
def corpus_anotar():
    """Fila de anotação: entrega a próxima notícia sem rótulo."""
    con = conexao()
    pendentes = corpus.listar(con, {"rotulada": "nao"}, 1)
    if not pendentes:
        flash("Nenhuma notícia sem rótulo. O corpus está todo anotado.", "ok")
        return redirect(url_for("corpus_repositorio"))
    return redirect(url_for("corpus_noticia", noticia_id=pendentes[0]["id"]))


@app.route("/corpus/<int:noticia_id>/excluir", methods=["POST"])
def corpus_excluir(noticia_id: int):
    con = conexao()
    with con:
        corpus.remover(con, noticia_id)
    flash(f"Notícia #{noticia_id} removida do corpus.", "ok")
    return redirect(url_for("corpus_repositorio"))


@app.route("/corpus/dividir", methods=["POST"])
def corpus_dividir():
    con = conexao()
    try:
        semente = int(request.form.get("semente", 42))
    except ValueError:
        semente = 42
    with con:
        contagem = corpus.dividir(con, semente=semente)
    flash(
        f"Divisão estratificada com semente {semente}: {contagem['treino']} treino,"
        f" {contagem['validacao']} validação, {contagem['teste']} teste.",
        "ok",
    )
    return redirect(url_for("corpus_repositorio"))


COLUNAS_CORPUS = [
    ("id", "ID"), ("link", "Link"), ("texto", "Texto"),
    ("indicadores", "Indicadores"), ("rotulada", "Anotada"),
    ("particao", "Conjunto"), ("anotador", "Anotador"), ("criado_em", "Criado em"),
]


@app.route("/corpus/exportar")
def corpus_exportar():
    """O corpus precisa sair do banco para ser anexado ao TCC e reproduzido."""
    from openpyxl import Workbook

    planilha = Workbook()
    planilha.remove(planilha.active)
    aba = planilha.create_sheet("corpus")
    aba.append([rotulo for _, rotulo in COLUNAS_CORPUS])
    for registro in corpus.listar(conexao(), _filtros_de_corpus()):
        aba.append([
            ", ".join(registro["indicadores"]) if campo == "indicadores"
            else registro.get(campo)
            for campo, _ in COLUNAS_CORPUS
        ])
    return _enviar_planilha(planilha, "corpus-treino")


@app.route("/corpus/importar", methods=["POST"])
def corpus_importar():
    """Importa um lote de notícias de planilha, para quando o corpus já existe
    fora da aplicação. Colunas aceitas: link, texto (obrigatória) e indicadores
    (nomes separados por vírgula; vazio manda a notícia para a fila)."""
    from openpyxl import load_workbook

    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        flash("Selecione uma planilha .xlsx para importar.", "erro")
        return redirect(url_for("corpus_repositorio"))

    try:
        planilha = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as erro:
        flash(f"Não foi possível ler a planilha: {erro}", "erro")
        return redirect(url_for("corpus_repositorio"))

    aba = planilha[planilha.sheetnames[0]]
    linhas = aba.iter_rows(values_only=True)
    cabecalho = [str(c or "").strip().lower() for c in next(linhas, [])]
    if "texto" not in cabecalho:
        flash("A planilha precisa ter uma coluna chamada 'texto'.", "erro")
        return redirect(url_for("corpus_repositorio"))

    aprovar_importadas = bool(request.form.get("aprovar"))
    con = conexao()
    importadas = ignoradas = 0
    with con:
        for linha in linhas:
            registro = dict(zip(cabecalho, linha))
            texto = str(registro.get("texto") or "").strip()
            if not texto:
                ignoradas += 1
                continue
            bruto = registro.get("indicadores") or ""
            corpus.adicionar(con, {
                "link": registro.get("link"),
                "texto": texto,
                "indicadores": [p.strip() for p in str(bruto).split(",") if p.strip()],
                "rotulada": registro.get("anotada"),
                # Planilha não vem aprovada por padrão: ninguém aqui revisou o
                # rótulo dela. A caixa existe para o corpus que já foi revisado
                # fora daqui, que é o caso de uso real da importação.
                "aprovada": "1" if aprovar_importadas else "0",
            }, analista_atual())
            importadas += 1

    flash(
        f"{importadas} notícia(s) importada(s)"
        + (" e aprovadas para o treino." if aprovar_importadas else
           ", esperando revisão — elas não treinam o modelo até serem aprovadas.")
        + (f" {ignoradas} linha(s) sem texto ignorada(s)." if ignoradas else ""),
        "ok",
    )
    return redirect(url_for("corpus_repositorio"))


@app.route("/modelos")
def modelos():
    """Resultados do treino: a tabela comparativa e as matrizes de confusão."""
    con = conexao()
    execucoes = {}
    for tarefa in corpus.TAREFAS:
        linhas = con.execute(
            "SELECT * FROM treino_execucao WHERE tarefa = ? ORDER BY id DESC",
            (tarefa,),
        ).fetchall()
        registros = []
        for linha in linhas:
            registro = {k: linha[k] for k in linha.keys()}
            registro["hiperparametros"] = db.json_carregar(linha["hiperparametros"], {})
            registro["metricas"] = db.json_carregar(linha["metricas"], {})
            registro["rotulos"] = db.json_carregar(linha["rotulos"], [])
            registro["limiares"] = db.json_carregar(linha["limiares"], {})
            registro["historico"] = db.json_carregar(linha["historico"], [])
            registros.append(registro)
        execucoes[tarefa] = registros

    return render_template(
        "modelos.html",
        execucoes=execucoes,
        # Uma linha por modelo, com a melhor configuração dele: é o que responde
        # "qual modelo está ganhando" sem obrigar a ler as 24 execuções.
        placar={t: treinamento.placar(r) for t, r in execucoes.items()},
        tarefas=corpus.TAREFAS,
        catalogo=treinamento.MODELOS,
        grade=treinamento.GRADE,
        estrategias=treinamento.ESTRATEGIAS,
        ambiente=treinamento.descrever_ambiente(),
        estatisticas=corpus.estatisticas(con),
        treino=treino_em_curso.situacao(),
        treinados={t: inferencia.treinado(t) for t in corpus.TAREFAS},
        # O que está em produção, com que estratégia e com que cortes: a tela de
        # modelos é onde se compara, e sem isto não se saberia qual linha da
        # tabela é a que a aplicação está usando de fato.
        em_producao={t: inferencia.resumo(t) or {} for t in corpus.TAREFAS},
        # Se o catálogo ganhou indicador depois do último treino. Sem este
        # aviso, o indicador novo existia na tela e no prompt e não existia
        # para o classificador, sem nada dizer.
        defasagem=inferencia.defasagem(),
        rotulos=indicadores.rotulos(),
    )


@app.route("/modelos/treinar", methods=["POST"])
def treinar_modelo():
    """Dispara o treino em segundo plano.

    O treino NÃO roda dentro desta requisição: ele leva minutos e carrega um
    BERT inteiro na memória. Aqui só se lança o processo e se guarda o PID; a
    tela acompanha pelo registro.
    """
    con = conexao()

    # Distribui as pendentes ANTES de julgar a prontidão. Sem isto, um corpus
    # recém-rotulado seria recusado com "a divisão ainda não foi sorteada" — e
    # a distribuição é justamente o que falta, não um impedimento. Só o que
    # está sem partição é tocado; o que já foi sorteado continua onde está.
    pendentes = corpus.pendentes_de_particao(con)
    if pendentes:
        with con:
            corpus.dividir(con)

    estatisticas = corpus.estatisticas(con)
    prontidao = estatisticas["prontidao"]["indicadores"]
    if not prontidao["pode_treinar"]:
        recado = f"Ainda não dá para treinar: {prontidao['motivo']}."
        esperando = estatisticas["rotuladas_sem_aprovar"]
        if esperando:
            # A causa mais provável quando há rótulo e o treino recusa: o
            # rótulo existe e ninguém o aprovou. Dizer "rotule mais notícias"
            # nesse caso manda a pessoa fazer o trabalho errado.
            recado += (
                f" Há {esperando} notícia(s) rotuladas esperando revisão —"
                " aprovar o que já está certo pode ser o bastante."
            )
        else:
            recado += " Rotule mais notícias na fila de recebidas ou no corpus."
        flash(recado, "erro")
        return redirect(url_for("modelos"))

    resultado = treino_em_curso.iniciar({
        "estrategia": request.form.get("estrategia") or None,
        "modelos": request.form.getlist("modelos"),
        "grade": request.form.get("grade") or None,
        "dispositivo": request.form.get("dispositivo") or None,
        "dividir": bool(request.form.get("dividir")),
    })
    if not resultado["iniciado"]:
        flash(resultado["motivo"], "aviso")
    else:
        if pendentes:
            flash(
                f"{pendentes} notícia(s) sem partição foram distribuídas em"
                " treino/validação/teste antes de começar.",
                "ok",
            )
        flash(
            f"Treino iniciado em segundo plano (processo {resultado['pid']})."
            " Pode fechar esta página: o resultado aparece aqui quando"
            " terminar. Se houver aviso na prontidão acima, ele também vale"
            " para este treino.",
            "ok",
        )
    return redirect(url_for("modelos"))


@app.route("/modelos/<int:execucao_id>")
def modelo_execucao(execucao_id: int):
    linha = conexao().execute(
        "SELECT * FROM treino_execucao WHERE id = ?", (execucao_id,)
    ).fetchone()
    if linha is None:
        flash("Execução não encontrada.", "erro")
        return redirect(url_for("modelos"))

    execucao = {k: linha[k] for k in linha.keys()}
    execucao["hiperparametros"] = db.json_carregar(linha["hiperparametros"], {})
    execucao["metricas"] = db.json_carregar(linha["metricas"], {})
    execucao["rotulos"] = db.json_carregar(linha["rotulos"], [])
    execucao["limiares"] = db.json_carregar(linha["limiares"], {})
    execucao["historico"] = db.json_carregar(linha["historico"], [])
    return render_template(
        "modelo_execucao.html", execucao=execucao, catalogo=treinamento.MODELOS,
        estrategias=treinamento.ESTRATEGIAS,
    )


@app.route("/testar", methods=["GET", "POST"])
def testar():
    """Playground: o que o BERT responde, antes de qualquer LLM, e quais
    prompts de indicador essa resposta faria a LLM carregar."""
    texto = request.form.get("texto", "") if request.method == "POST" else ""
    resultado = erro = None

    # Vazio significa "use o limiar calibrado de cada indicador", que é o
    # comportamento de produção. Um número substitui todos os cortes por ele e
    # serve só para ver o que apareceria ou desapareceria com outro corte.
    limiar_pedido = (request.form.get("limiar") or "").strip()
    try:
        limiar = float(limiar_pedido.replace(",", ".")) if limiar_pedido else None
    except ValueError:
        limiar = None

    if request.method == "POST" and texto.strip():
        try:
            resultado = inferencia.prever(texto, limiar=limiar)
            resultado["modelo"] = str(inferencia.caminho_do_modelo())
        except RuntimeError as falha:
            # Sem modelo treinado não há predição — mas o texto não se perde:
            # ele é gravado logo abaixo com `classificado = 0`, à espera do
            # primeiro treino.
            erro = str(falha)

        con = conexao()
        with con:
            # A entrada é gravada SEMPRE, com predição ou sem ela.
            entrada_id = auditoria.registrar_entrada(
                con, texto, resultado,
                link=(request.form.get("link") or "").strip() or None,
                origem="playground",
            )
            # O teste também entra no banco de classificações: um erro visto no
            # playground é tão útil para o próximo treino quanto um de produção.
            if resultado is not None:
                auditoria.registrar(
                    con, resultado, texto, origem="playground",
                    entrada_id=entrada_id,
                )

    prompt = None
    if resultado is not None:
        # Mostrar o prompt que a marcação geraria é o que torna visível a
        # economia: sem indicador marcado, a LLM recebe algumas linhas; com
        # três, recebe três blocos, e nunca os do catálogo inteiro.
        prompt = {
            "nucleo": len(classificador.montar_nucleo()),
            "blocos": classificador.montar_prompt_indicadores(resultado["previsto"]),
            "completo": len(classificador.montar_prompt_indicadores(None)),
        }

    return render_template(
        "testar.html",
        texto=texto,
        resultado=resultado,
        erro=erro,
        prompt=prompt,
        limiar_pedido=limiar_pedido,
        treinado=inferencia.treinado(),
        estrategia=inferencia.estrategia(),
        limiares=inferencia.limiares(),
        rotulos=indicadores.rotulos(),
    )


# --------------------------------------------------------------------------
# Banco das classificações do BERT
# --------------------------------------------------------------------------

def _filtros_de_classificacao() -> dict:
    return {
        chave: request.args.get(chave, "").strip()
        for chave in ("revisado", "classe", "sem_classe", "origem", "busca")
    }


@app.route("/classificacoes")
def classificacoes():
    """O que o BERT marcou, para conferir depois o que ele errou."""
    con = conexao()
    filtros = _filtros_de_classificacao()
    total = auditoria.contar(con, filtros)

    paginas = max(1, -(-total // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1

    return render_template(
        "classificacoes.html",
        registros=auditoria.listar(
            con, filtros, POR_PAGINA, (pagina - 1) * POR_PAGINA
        ),
        filtros=filtros,
        filtros_ativos={c: v for c, v in filtros.items() if v},
        estatisticas=auditoria.estatisticas(con),
        # O banco de entradas: todo texto que entrou, inclusive o que não
        # recebeu indicador nenhum e o que chegou antes de haver modelo.
        entradas=auditoria.estatisticas_de_entrada(con),
        catalogo=indicadores.do_bert(),
        origem_modelo=inferencia.origem(),
        total=total,
        pagina=pagina,
        paginas=paginas,
    )


@app.route("/classificacoes/<int:registro_id>/revisar", methods=["POST"])
def revisar_classificacao(registro_id: int):
    con = conexao()
    with con:
        auditoria.revisar(
            con, registro_id, request.form.getlist("classes"), analista_atual(),
            request.form.get("observacao", ""),
        )
        if request.form.get("ao_corpus"):
            novo = auditoria.enviar_ao_corpus(con, registro_id, analista_atual())
            flash(
                f"Revisão salva e enviada ao corpus como notícia #{novo}."
                " É assim que o erro do modelo vira exemplo do próximo treino.",
                "ok",
            )
            return redirect(url_for("corpus_noticia", noticia_id=novo))
    flash("Revisão salva.", "ok")
    return redirect(request.referrer or url_for("classificacoes"))


@app.route("/classificacoes/<int:registro_id>/corpus", methods=["POST"])
def classificacao_ao_corpus(registro_id: int):
    con = conexao()
    with con:
        novo = auditoria.enviar_ao_corpus(con, registro_id, analista_atual())
    flash(f"Notícia #{novo} criada no corpus com estes rótulos.", "ok")
    return redirect(url_for("corpus_noticia", noticia_id=novo))


@app.route("/recebidas")
def fila_recebidas():
    """A fila de entrada: notícia que chegou e ainda não virou nada.

    É a primeira tela do fluxo novo. Sem modelo treinado, tudo cai aqui para o
    analista rotular — e é essa rotulagem que constrói o corpus do primeiro
    treino.
    """
    con = conexao()
    filtros = {
        "estado": request.args.get("estado", ""),
        "busca": request.args.get("busca", "").strip(),
    }
    por_pagina = 20
    total = recebidas.contar(con, filtros)
    paginas = max(1, -(-total // por_pagina))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    return render_template(
        "recebidas.html",
        lista=recebidas.listar(con, filtros, por_pagina, (pagina - 1) * por_pagina),
        estatisticas=recebidas.estatisticas(con),
        estados=recebidas.ESTADOS,
        classes=indicadores.do_bert(),
        # Aqui o aviso importa mais que em /modelos: é esta tela que decide a
        # faixa da notícia, e um indicador que o modelo não conhece faz a
        # notícia ser arquivada como se não tivesse nada.
        defasagem=inferencia.defasagem(),
        filtros=filtros,
        pagina=pagina,
        paginas=paginas,
        total=total,
    )


@app.route("/recebidas/nova", methods=["POST"])
def receber_noticia():
    texto = request.form.get("texto", "").strip()
    if not texto:
        flash("O texto da notícia é obrigatório.", "erro")
        return redirect(url_for("fila_recebidas"))

    con = conexao()
    with con:
        resultado = recebidas.adicionar(
            con,
            {
                "link": request.form.get("link"),
                "data_noticia": request.form.get("data_noticia"),
                "texto": texto,
            },
            analista_atual(),
        )
    if resultado["repetida"]:
        flash(f"Esta notícia já está na fila (#{resultado['id']}).", "aviso")
    elif resultado["marcados"]:
        flash(
            f"Notícia #{resultado['id']} recebida. O classificador marcou:"
            f" {', '.join(resultado['marcados'])}.",
            "ok",
        )
    else:
        flash(
            f"Notícia #{resultado['id']} recebida e à espera de rotulagem."
            " Sem modelo treinado, quem decide é você — e é dessa decisão que"
            " o primeiro modelo vai aprender.",
            "ok",
        )
    return redirect(url_for("fila_recebidas"))


@app.route("/recebidas/<int:recebida_id>/rotular", methods=["POST"])
def rotular_recebida(recebida_id: int):
    con = conexao()
    with con:
        resultado = recebidas.rotular(
            con, recebida_id,
            request.form.getlist("indicadores"),
            analista_atual(),
        )
    if resultado["indicadores"]:
        flash(
            f"Rotulada com {', '.join(resultado['indicadores'])} e enviada ao"
            f" corpus (#{resultado['corpus_id']}). Já pode virar ocorrência.",
            "ok",
        )
    else:
        flash(
            f"Guardada como exemplo SEM fenômeno (corpus #{resultado['corpus_id']})."
            " Não vira ocorrência — e é esta notícia que ensina o classificador"
            " a recusar o que está fora do escopo.",
            "ok",
        )
    return redirect(url_for("fila_recebidas"))


@app.route("/recebidas/reprocessar", methods=["POST"])
def reprocessar_arquivo():
    """Passa o classificador de novo sobre o acervo.

    É a razão de o arquivo existir: treinar um modelo novo ou acrescentar um
    indicador muda a resposta para notícias que já estão lá, e sem este passo
    elas ficariam com o veredito de um modelo que não existe mais.
    """
    con = conexao()
    if not inferencia.treinado():
        flash(
            "Não há modelo treinado: não há o que reavaliar. O arquivo existe"
            " justamente para esperar por ele.",
            "aviso",
        )
        return redirect(url_for("fila_recebidas", estado="arquivo"))
    with con:
        r = recebidas.reprocessar(con)
    if r["movidas"]:
        mudou = ", ".join(f"{n} para {recebidas.ESTADOS[e]}" for e, n in r["movidas"].items())
        flash(f"{r['lidas']} notícia(s) relidas. Mudaram: {mudou}.", "ok")
    else:
        flash(
            f"{r['lidas']} notícia(s) relidas e nenhuma mudou de faixa:"
            " o modelo atual continua sem achar indicador nelas.",
            "aviso",
        )
    return redirect(url_for("fila_recebidas", estado="arquivo"))


@app.route("/recebidas/<int:recebida_id>/excluir", methods=["POST"])
def excluir_recebida(recebida_id: int):
    con = conexao()
    with con:
        con.execute("DELETE FROM recebida WHERE id = ?", (recebida_id,))
    flash("Notícia removida da fila.", "ok")
    return redirect(url_for("fila_recebidas"))


@app.route("/locais")
def locais_e_listas():
    """Cadastro de locais e das listas fechadas que o analista mantém.

    Duas coisas na mesma tela porque é o mesmo tipo de trabalho: manter a lista
    de valores que os formulários oferecem. Separá-las em duas telas faria
    procurar em dois lugares pela mesma razão.
    """
    sigla = (request.args.get("estado") or "").strip().upper()
    cidade = (request.args.get("cidade") or "").strip()
    if sigla and cidade and not mod_locais.existe_cidade(sigla, cidade):
        cidade = ""
    return render_template(
        "locais.html",
        estados=mod_locais.estados(),
        sigla=sigla,
        cidade=cidade,
        cidades=mod_locais.cidades(sigla) if sigla else [],
        bairros=mod_locais.bairros(sigla, cidade) if sigla and cidade else [],
        estatisticas=mod_locais.estatisticas(),
        editaveis={
            nome: {"rotulo": rotulo, "itens": getattr(vocabulario, nome)}
            for nome, rotulo in vocabulario.EDITAVEIS.items()
        },
    )


@app.route("/locais/cidade", methods=["POST"])
def locais_cidade():
    sigla = (request.form.get("estado") or "").strip().upper()
    acao = request.form.get("acao")
    nome = request.form.get("cidade", "")
    recado = (
        mod_locais.remover_cidade(sigla, nome) if acao == "remover"
        else mod_locais.acrescentar_cidade(sigla, nome)
    )
    flash(recado, "ok" if "já está" not in recado and "Informe" not in recado else "aviso")
    return redirect(url_for("locais_e_listas", estado=sigla))


@app.route("/locais/bairro", methods=["POST"])
def locais_bairro():
    sigla = (request.form.get("estado") or "").strip().upper()
    cidade = request.form.get("cidade", "")
    acao = request.form.get("acao")
    nome = request.form.get("bairro", "")
    recado = (
        mod_locais.remover_bairro(sigla, cidade, nome) if acao == "remover"
        else mod_locais.acrescentar_bairro(sigla, cidade, nome)
    )
    flash(recado, "ok" if "já está" not in recado and "Informe" not in recado else "aviso")
    return redirect(url_for("locais_e_listas", estado=sigla, cidade=cidade))


@app.route("/locais/lista", methods=["POST"])
def locais_lista():
    """Salva uma lista fechada inteira, do jeito que a tela a devolveu.

    O VALOR de um item já gravado não pode mudar: ele é o que está nas
    ocorrências antigas, e reescrevê-lo deixaria aqueles registros apontando
    para um valor que não existe mais. Por isso o campo do valor é somente
    leitura nos itens existentes, e a tela só deixa editar o rótulo ou
    acrescentar item novo.
    """
    nome = request.form.get("lista", "")
    if nome not in vocabulario.EDITAVEIS:
        flash("Lista desconhecida.", "erro")
        return redirect(url_for("locais_e_listas"))

    valores = request.form.getlist("valor")
    rotulos = request.form.getlist("rotulo")
    manter = set(request.form.getlist("manter"))

    pares = []
    for valor, rotulo in zip(valores, rotulos):
        valor = (valor or "").strip()
        rotulo = (rotulo or "").strip()
        if not valor or not rotulo:
            continue
        # Desmarcar a caixa é o que remove: um botão de excluir por linha daria
        # uma gravação por item e perderia a ordem entre elas.
        if valores.count(valor) and valor not in manter:
            continue
        pares.append((valor, rotulo))

    novo_valor = (request.form.get("novo_valor") or "").strip()
    novo_rotulo = (request.form.get("novo_rotulo") or "").strip()
    if novo_valor and novo_rotulo:
        if any(v == novo_valor for v, _ in pares):
            flash(f"O valor {novo_valor} já existe nesta lista.", "aviso")
        else:
            pares.append((novo_valor, novo_rotulo))

    if not pares:
        flash("A lista ficaria vazia; nada foi salvo.", "erro")
        return redirect(url_for("locais_e_listas"))

    vocabulario.salvar_lista(nome, pares)
    flash(
        f"{vocabulario.EDITAVEIS[nome]}: {len(pares)} item(ns) salvos."
        " As telas e a validação já usam a lista nova.",
        "ok",
    )
    return redirect(url_for("locais_e_listas"))


@app.route("/testar/salvar", methods=["POST"])
def testar_salvar():
    """Leva um texto do playground para o corpus. É assim que um erro visto no
    teste vira exemplo de treino em vez de só uma anotação perdida."""
    texto = request.form.get("texto", "").strip()
    if not texto:
        flash("Nada para salvar.", "erro")
        return redirect(url_for("testar"))

    con = conexao()
    with con:
        novo = corpus.adicionar(con, {
            "link": request.form.get("link"),
            "texto": texto,
        }, analista_atual())
    flash(f"Notícia #{novo} enviada ao corpus. Agora defina os rótulos.", "ok")
    return redirect(url_for("corpus_noticia", noticia_id=novo))


if __name__ == "__main__":
    db.iniciar()
    app.run(debug=True, port=5000)
