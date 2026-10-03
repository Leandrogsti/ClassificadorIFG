"""Gera a documentação de TCC em docx a partir das fontes reais do projeto.

As tabelas de apêndice saem de categorias_fonte.json, indicadores_fonte.json e
do catálogo de modelos em treinamento.py, não de texto copiado à mão. Alterar
uma categoria e regerar mantém o documento em dia com o sistema que ele
descreve, em vez de deixar os dois divergirem em silêncio.

    python gerar_documentacao.py
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJETO = Path(__file__).parent
SAIDA = PROJETO / "docs"
SAIDA.mkdir(exist_ok=True)

import corpus as mod_corpus
import treinamento as mod_treino
import vocabulario as mod_vocab
import indicadores as mod_ind

CATEGORIAS = json.loads((PROJETO / "categorias_fonte.json").read_text(encoding="utf-8"))["categorias"]
INDICADORES = json.loads((PROJETO / "indicadores_fonte.json").read_text(encoding="utf-8"))["indicadores"]

CINZA = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# ---------------------------------------------------------------- formatacao
secao = doc.sections[0]
secao.page_width = Cm(21.0)
secao.page_height = Cm(29.7)
secao.left_margin = Cm(3)
secao.top_margin = Cm(3)
secao.right_margin = Cm(2)
secao.bottom_margin = Cm(2)

def fixar_fonte(estilo, nome="Times New Roman"):
    """Os estilos internos do Word apontam para a fonte do TEMA, e o atributo de
    tema vence o w:ascii que o python-docx escreve. Sem remover o tema, os
    titulos saem na fonte sem serifa do tema padrao."""
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for atributo in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        chave = qn("w:" + atributo)
        if chave in rfonts.attrib:
            del rfonts.attrib[chave]
    for atributo in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + atributo), nome)


normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
fixar_fonte(normal)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for nome_estilo in ("List Bullet", "Caption"):
    try:
        fixar_fonte(doc.styles[nome_estilo])
    except KeyError:
        pass

for nivel, tamanho in [(1, 14), (2, 13), (3, 12)]:
    estilo = doc.styles[f"Heading {nivel}"]
    estilo.font.name = "Times New Roman"
    fixar_fonte(estilo)
    estilo.font.size = Pt(tamanho)
    estilo.font.bold = True
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    estilo.paragraph_format.space_before = Pt(18)
    estilo.paragraph_format.space_after = Pt(10)
    estilo.paragraph_format.line_spacing = 1.5
    estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    estilo.paragraph_format.keep_with_next = True


def p(texto="", recuo=True, centro=False, italico=False, negrito=False, tamanho=None, espaco=6):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    if centro:
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif recuo:
        pf.first_line_indent = Cm(1.25)
    pf.space_after = Pt(espaco)
    if texto:
        run = par.add_run(texto)
        run.italic = italico
        run.bold = negrito
        if tamanho:
            run.font.size = Pt(tamanho)
    return par


def h(nivel, texto):
    return doc.add_heading(texto, level=nivel)


def bullet(texto, negrito_ate=None):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.line_spacing = 1.5
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.left_indent = Cm(1.25)
    if negrito_ate:
        par.add_run(negrito_ate).bold = True
        par.add_run(texto)
    else:
        par.add_run(texto)
    return par


def codigo(linhas):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(1.25)
    pf.space_before = Pt(6)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, linha in enumerate(linhas):
        if i:
            par.add_run().add_break()
        run = par.add_run(linha)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
    sombrear(par, "F2F3F5")
    return par


def sombrear(par_ou_celula, cor_hex):
    elemento = par_ou_celula._p if hasattr(par_ou_celula, "_p") else par_ou_celula._tc
    pr = elemento.get_or_add_pPr() if hasattr(par_ou_celula, "_p") else elemento.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    pr.append(shd)


def legenda(texto, acima=False):
    par = doc.add_paragraph()
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_before = Pt(10 if acima else 2)
    par.paragraph_format.space_after = Pt(4 if acima else 12)
    par.paragraph_format.line_spacing = 1.0
    run = par.add_run(texto)
    run.font.size = Pt(10)
    run.font.color.rgb = CINZA
    return par


def tabela(cabecalho, linhas, larguras_cm):
    tab = doc.add_table(rows=1, cols=len(cabecalho))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False

    def preencher(celula, texto, largura, negrito=False):
        celula.width = Cm(largura)
        celula.text = ""
        par = celula.paragraphs[0]
        pf = par.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(2)
        pf.space_before = Pt(2)
        # Celula estreita com texto justificado abre vaos enormes entre palavras.
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = par.add_run(str(texto))
        run.bold = negrito
        run.font.size = Pt(10)

    for i, (titulo, largura) in enumerate(zip(cabecalho, larguras_cm)):
        preencher(tab.rows[0].cells[i], titulo, largura, negrito=True)
        sombrear(tab.rows[0].cells[i], "E6EAF0")

    # Repete o cabecalho quando a tabela atravessa paginas.
    tr_pr = tab.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    tr_pr.append(repetir)

    for valores in linhas:
        celulas = tab.add_row().cells
        for i, valor in enumerate(valores):
            preencher(celulas[i], valor, larguras_cm[i])
    return tab


def quebra():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sumario():
    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    filho = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "Atualize o sumário no Word: clique com o botão direito e escolha “Atualizar campo”."
    filho.append(texto)
    fld.append(filho)
    par._p.append(fld)


# ---------------------------------------------------------------------- capa
for _ in range(2):
    p(centro=True)
p("INSTITUTO FEDERAL DE GOIÁS", centro=True, recuo=False, negrito=True, espaco=0)
p("TRABALHO DE CONCLUSÃO DE CURSO", centro=True, recuo=False, espaco=0)
for _ in range(4):
    p(centro=True)
p("CLASSIFICAÇÃO AUTOMÁTICA DE NOTÍCIAS DE VIOLÊNCIA ARMADA",
  centro=True, recuo=False, negrito=True, tamanho=16, espaco=0)
p("COM MODELOS BERT EM PORTUGUÊS E VALIDAÇÃO POR MODELO DE LINGUAGEM",
  centro=True, recuo=False, negrito=True, tamanho=16, espaco=0)
for _ in range(2):
    p(centro=True)
p("Documentação de método, implementação e avaliação",
  centro=True, recuo=False, italico=True, espaco=0)
for _ in range(8):
    p(centro=True)
p("Goiânia", centro=True, recuo=False, espaco=0)
p("2026", centro=True, recuo=False, espaco=0)

quebra()

# -------------------------------------------------------------------- resumo
h(1, "Resumo")
p("Este documento descreve o método, a implementação e o protocolo de avaliação "
  "de um sistema que classifica notícias de violência armada em indicadores, "
  "substituindo a triagem manual que hoje limita o volume de notícias analisadas "
  "e a velocidade de atualização dos indicadores. A classificação é organizada "
  "em duas etapas encadeadas. A primeira é binária e tem premissa única: "
  "determinar se a notícia relata disparo de arma de fogo. A segunda, aplicada "
  "somente ao que passa pela primeira, atribui a motivação principal entre "
  f"{len(CATEGORIAS)} categorias e os indicadores de contexto. Para cada etapa "
  f"são comparados {len(mod_treino.MODELOS)} modelos pré-treinados — monolíngues "
  "de português e multilíngues — sob uma grade de "
  f"{len(mod_treino.GRADE)} configurações de hiperparâmetros, com relato de "
  "acurácia, precisão, recall e medida F1, em versões macro, ponderada e por "
  "classe, acompanhadas da matriz de confusão. Uma camada de modelo de linguagem "
  "de grande porte opera após os classificadores, encarregada da confirmação "
  "contextual e da consolidação de vítimas. O documento detalha a construção do "
  "corpus anotado manualmente, o protocolo de divisão amostral, os critérios de "
  "seleção de hiperparâmetros e as ameaças à validade dos resultados.")
p()
p("Palavras-chave: processamento de linguagem natural; BERT; BERTimbau; "
  "classificação de texto; violência armada; indicadores de segurança pública.",
  recuo=False)

quebra()

# ------------------------------------------------------------------- sumario
h(1, "Sumário")
sumario()

quebra()

# --------------------------------------------------------------- introducao
h(1, "1 Introdução")

h(2, "1.1 Contextualização")
p("Instituições que monitoram violência armada constroem seus indicadores a "
  "partir da leitura de notícias. Um analista lê cada texto, decide se o caso "
  "pertence ao escopo, classifica a dinâmica do episódio e registra as vítimas. "
  "O procedimento é confiável quando conduzido por pessoas treinadas, mas o seu "
  "custo cresce linearmente com o volume de notícias, e é esse custo que fixa o "
  "teto do que se consegue monitorar.")
p("A consequência não é apenas de produtividade. Quando a triagem não acompanha "
  "o fluxo de publicações, os indicadores passam a descrever um recorte parcial "
  "e defasado da realidade, o que compromete o seu uso em decisões de política "
  "pública e em comunicação pública.")

h(2, "1.2 Problema")
p("O problema tratado é o da triagem e classificação automáticas de notícias "
  "para alimentar indicadores de violência armada, preservando a qualidade da "
  "classificação manual. A dificuldade está em três pontos. Primeiro, a decisão "
  "de escopo é sutil: muitas notícias mencionam armas de fogo sem relatar "
  "disparo. Segundo, as categorias de motivação não são mutuamente evidentes a "
  "partir de palavras isoladas, pois dependem da relação entre quem atirou, "
  "contra quem e por quê. Terceiro, a contagem de vítimas exige consolidar "
  "fontes diferentes sobre o mesmo fato sem somá-las indevidamente.")

h(2, "1.3 Objetivos")
p("O objetivo geral é projetar, implementar e avaliar um sistema de "
  "classificação automática de notícias de violência armada em dois estágios, "
  "com camada posterior de validação por modelo de linguagem.")
p("Os objetivos específicos são:", recuo=False)
bullet("construir um corpus de notícias anotado manualmente, com registro da "
       "fonte de cada exemplo, e disponibilizar ferramenta própria de anotação;")
bullet("comparar modelos pré-treinados de português e multilíngues na tarefa "
       "binária de identificação de disparo de arma de fogo;")
bullet("comparar os mesmos modelos na tarefa de atribuição de motivação e na "
       "atribuição multirrótulo de indicadores de contexto;")
bullet("conduzir busca de hiperparâmetros e relatar o seu efeito sobre o "
       "desempenho;")
bullet("relatar acurácia, precisão, recall, medida F1 e matriz de confusão para "
       "cada combinação de modelo e configuração;")
bullet("avaliar o ganho da camada de modelo de linguagem sobre a saída dos "
       "classificadores.")

h(2, "1.4 Justificativa")
p("A arquitetura em dois estágios com validação posterior não é um detalhe de "
  "implementação, e sim a resposta a uma restrição de custo. Submeter toda "
  "notícia coletada a um modelo de linguagem de grande porte é caro e lento. Um "
  "classificador pequeno, treinado especificamente para a decisão de escopo, "
  "descarta antecipadamente o que está fora dele, e a chamada custosa é "
  "reservada ao que sobrou. O desenho também torna o sistema auditável: "
  "separando a decisão de escopo da decisão de categoria, é possível medir cada "
  "uma isoladamente e identificar em qual estágio um erro se originou.")

quebra()

# --------------------------------------------------------- fundamentacao
h(1, "2 Fundamentação teórica")

h(2, "2.1 Representações contextuais e a arquitetura Transformer")
p("Modelos de linguagem baseados na arquitetura Transformer representam cada "
  "palavra em função do contexto em que ela aparece, por meio do mecanismo de "
  "autoatenção. A diferença em relação a representações estáticas é decisiva "
  "para o problema aqui tratado: em “a polícia apreendeu a arma” e “a polícia "
  "atirou com a arma”, a palavra “arma” ocupa papéis distintos, e só uma "
  "representação sensível ao contexto distingue os dois casos.")

h(2, "2.2 BERT e o pré-treinamento")
p("O BERT é treinado em duas tarefas não supervisionadas sobre grandes volumes "
  "de texto, entre elas a predição de palavras mascaradas, o que o obriga a "
  "modelar dependências em ambas as direções da sentença. O resultado é um "
  "conjunto de pesos que já codifica regularidades sintáticas e semânticas da "
  "língua e que pode ser especializado para uma tarefa específica com uma "
  "quantidade de exemplos anotados muito menor do que a exigida por um modelo "
  "treinado a partir do zero.")

h(2, "2.3 Modelos para português")
p("Modelos multilíngues cobrem dezenas de idiomas com um único vocabulário de "
  "subpalavras, o que reduz a capacidade dedicada a cada língua. Modelos "
  "monolíngues de português, treinados sobre corpora da língua, tendem a "
  "segmentar melhor o vocabulário e a capturar construções próprias do "
  "português brasileiro, incluindo o registro jornalístico e os termos do campo "
  "da segurança pública. A comparação entre as duas famílias é um dos eixos "
  "deste trabalho, e não uma premissa adotada de antemão.")

h(2, "2.4 Ajuste fino")
p("O ajuste fino acrescenta ao modelo pré-treinado uma camada de classificação "
  "inicializada aleatoriamente e atualiza todos os pesos com os exemplos "
  "anotados da tarefa. Corpora pequenos tornam o procedimento sensível: taxas de "
  "aprendizado elevadas destroem o conhecimento adquirido no pré-treinamento, e "
  "um número grande de épocas leva o modelo a memorizar o conjunto de treino em "
  "vez de generalizar. É por isso que a busca de hiperparâmetros descrita na "
  "seção 8 trabalha em uma faixa estreita e deliberadamente conservadora.")

h(2, "2.5 Métricas de classificação")
p("Seja, para uma classe, VP o número de verdadeiros positivos, FP o de falsos "
  "positivos, VN o de verdadeiros negativos e FN o de falsos negativos. "
  "Definem-se:")
p()
tabela(
    ["Métrica", "Definição", "O que responde"],
    [
        ["Acurácia", "(VP + VN) / (VP + VN + FP + FN)",
         "Que fração de todas as decisões está correta."],
        ["Precisão", "VP / (VP + FP)",
         "Entre os casos apontados como da classe, quantos realmente são."],
        ["Recall", "VP / (VP + FN)",
         "Entre os casos que são da classe, quantos o modelo encontrou."],
        ["Medida F1", "2 · (precisão · recall) / (precisão + recall)",
         "Média harmônica entre precisão e recall, penalizando desequilíbrio."],
    ],
    [2.6, 6.2, 7.2],
)
legenda("Tabela 1 — Métricas de classificação adotadas.")

p("A escolha da forma de agregação entre classes é metodologicamente relevante. "
  "A média macro calcula a métrica para cada classe e tira a média simples, "
  "atribuindo o mesmo peso a todas; a média ponderada pesa cada classe pelo seu "
  "número de exemplos. Em um corpus desbalanceado, a acurácia e as médias "
  "ponderadas podem ser altas enquanto as classes raras são sistematicamente "
  "erradas, porque acertar a classe majoritária já garante a maior parte do "
  "resultado. Por esse motivo, a medida F1 macro é adotada como métrica "
  "principal de comparação, e as demais são reportadas como complemento.")

h(2, "2.6 Matriz de confusão")
p("A matriz de confusão dispõe os rótulos verdadeiros nas linhas e as predições "
  "nas colunas. A diagonal principal contém os acertos; cada célula fora dela "
  "indica com qual classe o modelo confundiu a classe da linha. Seu valor "
  "diagnóstico está justamente fora da diagonal: ela mostra se os erros se "
  "concentram em pares de categorias semanticamente próximas, o que sugere "
  "revisão dos critérios de anotação, ou se estão dispersos, o que sugere "
  "insuficiência de exemplos.")
p("Para tarefas multirrótulo, em que uma notícia pode receber mais de um "
  "indicador simultaneamente, não existe matriz única. Reporta-se então uma "
  "matriz 2×2 por indicador, cada uma tratando a presença daquele indicador "
  "como problema binário independente.")

quebra()

# ------------------------------------------------------------- arquitetura
h(1, "3 Arquitetura da solução")

h(2, "3.1 Visão geral")
p("O sistema recebe a notícia, reduz o texto, aplica o classificador de escopo "
  "e, se a notícia estiver no escopo, aplica os classificadores de categoria. "
  "A saída dos classificadores é então submetida ao modelo de linguagem, "
  "responsável pela confirmação contextual e pela consolidação da lista de "
  "vítimas. O resultado é gravado com situação “não aprovado” e segue para "
  "revisão humana.")
p()
tabela(
    ["Etapa", "Natureza", "Descrição"],
    [
        ["1. Abertura", "humana", "Analista informa data e endereço da ocorrência."],
        ["2. Inclusão", "humana", "Analista cola link e texto de uma ou mais notícias."],
        ["3. Pré-processamento", "automática", "Descarte de trechos sem relação com a ocorrência."],
        ["4. Escopo", "automática", "Classificador binário: há disparo de arma de fogo?"],
        ["5. Categoria", "automática", "Motivação principal e indicadores de contexto."],
        ["6. Validação", "automática", "Modelo de linguagem confirma e consolida vítimas."],
        ["7. Verificação", "humana", "Analista confere e corrige a lista de vítimas."],
        ["8. Regras", "automática", "Indicadores dependentes das vítimas são calculados."],
        ["9. Registro", "automática", "Gravação com situação “não aprovado”."],
        ["10. Aprovação", "humana", "Responsável libera ou devolve o registro."],
    ],
    [3.4, 2.4, 10.2],
)
legenda("Tabela 2 — Etapas do fluxo, com indicação do agente responsável.")

h(2, "3.2 Por que duas etapas de classificação")
p("A separação entre escopo e categoria tem três justificativas. A primeira é "
  "de custo: a etapa de escopo é a que recebe o maior volume e, por ser "
  "binária, pode ser resolvida por um modelo pequeno com alta confiança. A "
  "segunda é de interpretabilidade: uma matriz de confusão binária distingue "
  "claramente os dois tipos de erro — admitir notícia fora do escopo e rejeitar "
  "notícia dentro dele — que têm consequências diferentes para o indicador "
  "final. A terceira é de qualidade do treino da segunda etapa: ao restringir a "
  "classificação de motivação às notícias que efetivamente relatam disparo, "
  "evita-se ensinar o modelo a atribuir motivação a textos para os quais a "
  "pergunta não se aplica.")

h(2, "3.3 O papel do modelo de linguagem")
p("A camada de modelo de linguagem não substitui os classificadores; ela atua "
  "sobre a saída deles. Sua função é confirmar o indicador com base em "
  "instruções textuais específicas e, principalmente, consolidar a lista de "
  "vítimas relendo em conjunto todas as notícias da mesma ocorrência. Essa "
  "releitura conjunta é o que impede que pessoas citadas em fontes distintas "
  "sejam contadas mais de uma vez.")
p("A existência do estágio anterior permite avaliar o ganho da camada custosa. "
  "Como a aplicação expõe a resposta dos classificadores isoladamente, é "
  "possível verificar se o modelo de linguagem está corrigindo o classificador "
  "ou apenas repetindo a sua decisão — pergunta que não teria resposta se as "
  "duas camadas fossem medidas apenas em conjunto.")

quebra()

# ------------------------------------------------------------------ corpus
h(1, "4 Passo 1: construção do corpus")

h(2, "4.1 A premissa da primeira tarefa")
p("A primeira tarefa tem uma premissa única e deliberadamente estreita: "
  "considera-se violência armada a notícia que relata disparo de arma de fogo. "
  "Essa definição é operacional, não conceitual — ela não afirma o que é "
  "violência armada em sentido amplo, e sim fixa o critério verificável que "
  "delimita o escopo do sistema.")
p("A consequência prática é que vários casos intuitivamente associados a armas "
  "ficam fora:", recuo=False)
bullet("ameaça com arma de fogo sem que tiro tenha sido efetuado;")
bullet("apreensão de armas e munições em operação, sem confronto;")
bullet("agressão, homicídio ou lesão por outro meio, como arma branca;")
bullet("porte ilegal de arma identificado em abordagem, sem disparo.")
p("Tratar esses casos como positivos tornaria a fronteira da classe difusa, e a "
  "métrica resultante mediria a inconsistência da anotação em vez da capacidade "
  "do modelo.")

h(2, "4.2 Necessidade de exemplos negativos")
p("Um corpus formado apenas por notícias de violência armada não é utilizável "
  "para a primeira tarefa. Sem exemplos negativos, o modelo não tem o que "
  "aprender a recusar, e a matriz de confusão perde as colunas que registram "
  "justamente o erro relevante — a admissão de notícia fora do escopo. A "
  "ferramenta de anotação, por isso, exibe a contagem das duas classes e impede "
  "o início do treino enquanto uma delas não tiver exemplos suficientes.")

h(2, "4.3 Anotação da segunda etapa")
p("A motivação principal é atribuída somente a notícias positivas na primeira "
  f"etapa, entre as {len(CATEGORIAS)} categorias do Apêndice A. A anotação segue "
  "dois princípios herdados do manual de curadoria: classifica-se o fato que "
  "originou os tiros, não o seu desfecho, e consideram-se apenas as "
  "circunstâncias diretamente ligadas ao momento do disparo, desprezando "
  "antecedentes e consequências que não constituam novo disparo.")
p("Os indicadores de contexto formam tarefa multirrótulo independente: uma "
  "mesma notícia pode receber mais de um, ou nenhum.")

h(2, "4.4 Rastreabilidade")
p("Cada exemplo do corpus guarda o endereço eletrônico e o texto integral da "
  "notícia que o originou, além da identificação do anotador, do instante da "
  "anotação e de um campo livre para a justificativa do rótulo em casos de "
  "fronteira. O registro da fonte não é conveniência: sem ele, nenhum resultado "
  "de treino é verificável por terceiros, e a afirmação de desempenho perde "
  "sustentação. O campo de justificativa cumpre função adicional, ao permitir "
  "revisitar decisões difíceis quando um erro sistemático aparece na matriz de "
  "confusão.")

h(2, "4.5 Ferramenta de anotação")
p("A anotação é feita em interface própria, integrada à aplicação, que cumpre o "
  "papel de ambientes especializados do gênero. A notícia é inserida com link, "
  "título e texto; a primeira decisão é binária e apresentada de forma "
  "destacada, com a premissa do disparo reafirmada na própria tela; as "
  "definições das categorias de motivação ficam disponíveis durante a anotação, "
  "o que reduz a dependência de memória do anotador. Notícias inseridas sem "
  "rótulo entram em uma fila, e há navegação direta para o próximo item "
  "pendente.")
p("A interface informa, por tarefa, se o corpus já sustenta uma medição "
  f"honesta. O limiar adotado é de {mod_corpus.MINIMO_POR_CLASSE} exemplos por "
  "classe para considerar a métrica estável, e de três exemplos como mínimo "
  "absoluto para que a classe possa ser representada nas três partições. O "
  "critério é informativo: abaixo de uma dezena de exemplos no conjunto de "
  "teste, a variação de um único acerto desloca o recall em dezenas de pontos "
  "percentuais, e relatar tal número como resultado seria enganoso.")

h(2, "4.6 Dimensionamento esperado")
p("Para a tarefa binária, um corpus equilibrado com algumas centenas de "
  "exemplos em cada classe é suficiente para estimativas estáveis. Para a "
  f"tarefa de motivação, com {len(CATEGORIAS)} categorias, o requisito é "
  "substancialmente maior: mantendo a ordem de grandeza de algumas dezenas de "
  "exemplos por categoria no conjunto de treino, chega-se à casa do milhar de "
  "notícias anotadas. Categorias raras na realidade — e portanto raras no "
  "corpus — devem ser reportadas com o seu suporte explícito, e a discussão dos "
  "resultados precisa distinguir limitação do modelo de insuficiência amostral.")

quebra()

# --------------------------------------------------------- preprocessamento
h(1, "5 Passo 2: pré-processamento")
p("Textos jornalísticos coletados da web carregam material alheio ao fato: "
  "chamadas para outras matérias, créditos de fotografia, elementos de "
  "navegação e publicidade. Esse material consome espaço da janela de entrada do "
  "modelo sem contribuir para a decisão.")
p("O pré-processamento adotado opera no nível da sentença. Um modelo de "
  "processamento de linguagem natural segmenta o texto e preserva as sentenças "
  "que contenham pelo menos um termo do vocabulário relevante — extraído "
  "automaticamente das próprias definições das categorias — ou uma entidade "
  "nomeada de pessoa, local ou organização. Se o filtro for agressivo demais e "
  "restar menos de um mínimo de sentenças, o texto original é mantido, de modo "
  "que o procedimento nunca piore a entrada.")
p("Uma segunda forma de compressão, no nível da palavra, com remoção de "
  "palavras funcionais e lematização, foi implementada e deliberadamente "
  "deixada fora do fluxo padrão. Ela remove preposições como “por” e “contra”, "
  "que são exatamente os elementos que marcam quem atirou contra quem. A "
  "supressão dessas marcas inverteu a direção atribuída ao disparo em testes "
  "conduzidos durante o desenvolvimento, produzindo classificação incorreta. O "
  "caso é registrado aqui porque ilustra um risco geral: técnicas de "
  "normalização avaliadas apenas pela economia de tokens podem destruir a "
  "informação que sustenta a tarefa.")

quebra()

# ------------------------------------------------------------------ divisao
h(1, "6 Passo 3: divisão amostral")

h(2, "6.1 Partições e proporções")
p("O corpus é dividido em três partições: treino, usada para atualizar os "
  "pesos; validação, usada para escolher hiperparâmetros e observar "
  "sobreajuste; e teste, usada uma única vez por configuração, para produzir as "
  "métricas reportadas. A proporção de referência é de "
  f"{int(mod_corpus.PROPORCAO_PADRAO[0]*100)}%, "
  f"{int(mod_corpus.PROPORCAO_PADRAO[1]*100)}% e "
  f"{int(mod_corpus.PROPORCAO_PADRAO[2]*100)}%, respectivamente.")

h(2, "6.2 Estratificação")
p("A divisão é estratificada para que a proporção entre classes se mantenha nas "
  "três partições. A implementação adotada estratifica de forma exata pelo "
  "rótulo binário e distribui as motivações dentro de cada grupo por "
  "intercalação: as listas de cada motivação são embaralhadas e percorridas em "
  "paralelo, formando uma sequência em que o primeiro exemplo de cada categoria "
  "aparece antes do segundo de qualquer outra.")
p("A escolha merece registro porque uma solução aparentemente mais correta "
  "falha. Estratificar separadamente por cada par de rótulo binário e motivação "
  "produz, com quinze categorias, um grande número de estratos com um ou dois "
  "exemplos. Como um estrato pequeno não pode ser repartido em três, ele é "
  "atribuído integralmente ao treino, e o efeito agregado é esvaziar validação e "
  "teste da classe positiva. A intercalação preserva a proporção binária exata, "
  "espalha as motivações entre as partições e garante que nenhuma categoria "
  "fique sem exemplo de treino, por ordenar a sequência começando pelo treino.")

h(2, "6.3 Fixação da divisão")
p("A partição de cada exemplo é gravada em banco de dados, e não sorteada a "
  "cada execução. A razão é de comparabilidade: se o conjunto de teste mudasse "
  "entre execuções, as métricas de dois modelos passariam a ser medidas sobre "
  "amostras diferentes, e a diferença observada entre eles confundiria efeito do "
  "modelo com efeito da amostra. O sorteio usa semente explícita, registrada "
  "junto ao resultado, e é refeito somente quando o corpus cresce.")

quebra()

# ------------------------------------------------------------------ modelos
h(1, "7 Passo 4: modelos comparados")
p("A comparação abrange seis modelos pré-treinados, escolhidos para cobrir três "
  "eixos de variação: a oposição entre modelos monolíngues de português e "
  "multilíngues, o efeito do tamanho do modelo e o efeito da destilação.")
p()
tabela(
    ["Identificação", "Modelo", "Parâmetros", "Família"],
    [[chave, m["rotulo"], m["parametros"], m["tipo"]]
     for chave, m in mod_treino.MODELOS.items()],
    [3.6, 4.0, 2.4, 6.0],
)
legenda("Tabela 3 — Modelos submetidos à comparação.")

p("A inclusão de duas variantes de tamanho da mesma família permite separar o "
  "ganho atribuível à escala do ganho atribuível ao idioma do pré-treinamento. "
  "A variante destilada fornece o ponto de comparação de custo: se o seu "
  "desempenho for próximo ao das variantes completas, a escolha de "
  "implementação passa a ser decidida por tempo de inferência, e não por "
  "qualidade.")
p("Modelos cujo treino não se complete — por exemplo, por insuficiência de "
  "memória da unidade de processamento gráfico nas configurações de maior "
  "comprimento de sequência — têm a falha registrada com a respectiva mensagem. "
  "A informação integra o resultado: um modelo que não cabe no equipamento "
  "disponível é um achado de viabilidade, e omiti-lo da tabela daria a entender "
  "que não foi avaliado.")

quebra()

# ----------------------------------------------------------- hiperparametros
h(1, "8 Passo 5: busca de hiperparâmetros")

h(2, "8.1 Grade")
p("Cada modelo é treinado sob todas as configurações da grade abaixo, "
  f"totalizando {len(mod_treino.MODELOS) * len(mod_treino.GRADE)} execuções por "
  "tarefa.")
p()
tabela(
    ["#", "Taxa de aprendizado", "Lote", "Épocas", "Decaimento", "Comprimento"],
    [[i + 1, f"{h_['learning_rate']:.0e}", h_["batch_size"], h_["epocas"],
      h_["weight_decay"], h_["max_length"]]
     for i, h_ in enumerate(mod_treino.GRADE)],
    [1.0, 4.0, 1.6, 1.8, 3.0, 4.6],
)
legenda("Tabela 4 — Grade de hiperparâmetros.")

p("Os valores situam-se na faixa recomendada pela literatura de ajuste fino de "
  "modelos BERT, com taxas de aprendizado entre 2·10⁻⁵ e 5·10⁻⁵ e número "
  "reduzido de épocas. A grade é construída de modo que as configurações com "
  "mais épocas empreguem taxa menor, pois corpora pequenos sobreajustam "
  "rapidamente. Ao treino aplica-se ainda aquecimento linear da taxa de "
  "aprendizado em dez por cento dos passos e limitação da norma do gradiente, "
  "medidas usuais para estabilizar as primeiras atualizações.")

h(2, "8.2 Critério de seleção")
p("A configuração vencedora de cada tarefa é escolhida pela medida F1 macro "
  "obtida no conjunto de validação, nunca no de teste. A distinção é central "
  "para a validade do resultado: o conjunto de teste estima o desempenho em "
  "dados não vistos, e essa estimativa deixa de ser válida se ele for consultado "
  "para tomar decisões de modelagem. Escolher a configuração pelo teste "
  "transformaria o número reportado em um máximo sobre a grade, sistematicamente "
  "otimista.")
p("Quando uma tarefa não dispõe de conjunto de validação — situação possível "
  "com classes muito raras —, o critério recai sobre o teste, e essa condição é "
  "registrada explicitamente junto ao resultado, por deixar de ser uma "
  "estimativa limpa.")

h(2, "8.3 Acompanhamento do aprendizado")
p("Registram-se, por época, a perda média no conjunto de treino e a medida F1 "
  "macro no conjunto de validação. A leitura conjunta das duas séries é o que "
  "permite diagnosticar sobreajuste: perda decrescente acompanhada de medida F1 "
  "de validação estável ou em queda indica que o modelo está memorizando o "
  "treino em vez de generalizar, e sugere reduzir épocas ou taxa de aprendizado.")

quebra()

# ---------------------------------------------------------------- avaliacao
h(1, "9 Passo 6: protocolo de avaliação")

h(2, "9.1 Métricas reportadas")
p("Para cada execução são registradas a acurácia; a precisão, o recall e a "
  "medida F1 em média macro e em média ponderada pelo suporte; as mesmas três "
  "métricas por classe, acompanhadas do número de exemplos daquela classe no "
  "conjunto de teste; e a matriz de confusão completa. Para a tarefa "
  "multirrótulo reportam-se, adicionalmente, as médias micro e a acurácia de "
  "subconjunto, que exige acerto simultâneo de todos os rótulos de uma notícia.")
p("O suporte por classe é parte obrigatória do relato. Uma classe sem exemplos "
  "no conjunto de teste recebe, por convenção, medida F1 igual a zero, o que "
  "não representa erro do modelo e sim ausência de medição. Sem o suporte ao "
  "lado, o valor é lido incorretamente.")

h(2, "9.2 Estrutura da matriz de confusão")
p("A tabela a seguir exemplifica a disposição adotada para a tarefa binária, "
  "com os quatro resultados possíveis e a sua interpretação no domínio.")
p()
tabela(
    ["", "Predito: sem disparo", "Predito: com disparo"],
    [
        ["Real: sem disparo",
         "Verdadeiro negativo — notícia fora do escopo corretamente descartada.",
         "Falso positivo — notícia fora do escopo admitida; gera chamada de modelo de linguagem desnecessária e pode contaminar o indicador."],
        ["Real: com disparo",
         "Falso negativo — ocorrência real descartada; o caso simplesmente não entra na base e o indicador é subestimado.",
         "Verdadeiro positivo — ocorrência corretamente encaminhada à segunda etapa."],
    ],
    [3.4, 6.3, 6.3],
)
legenda("Tabela 5 — Interpretação da matriz de confusão da tarefa de escopo.")

p("A assimetria entre os dois tipos de erro é relevante para a escolha "
  "operacional do ponto de corte. O falso positivo tem custo de processamento e "
  "é passível de correção na etapa seguinte, inclusive pela revisão humana. O "
  "falso negativo é mais grave, porque a ocorrência não chega a existir no "
  "sistema e o erro não deixa rastro. Essa consideração favorece, para a "
  "primeira etapa, a priorização do recall da classe positiva.")

h(2, "9.3 Avaliação interativa")
p("Além das métricas agregadas, a aplicação disponibiliza interface de teste "
  "que recebe um texto e exibe a resposta do classificador com a distribuição "
  "de confiança sobre todas as classes, antes de qualquer intervenção do modelo "
  "de linguagem. A informação de confiança é parte do resultado: um acerto com "
  "probabilidade próxima ao limiar e outro com probabilidade alta descrevem "
  "situações distintas quanto à decisão de confiar no filtro ou encaminhar o "
  "caso à camada seguinte.")
p("A mesma interface permite enviar ao corpus um texto em que o modelo errou, "
  "fechando o ciclo entre avaliação e ampliação do conjunto anotado.")

quebra()

# ---------------------------------------------------------------------- llm
h(1, "10 Passo 7: validação por modelo de linguagem")
p("A camada final submete as notícias já filtradas a um modelo de linguagem de "
  "grande porte, com instruções textuais específicas por indicador. A saída do "
  "classificador é fornecida como sugestão, explicitamente sujeita a "
  "confirmação ou descarte, e não como decisão.")
p("A atribuição central dessa camada é a consolidação de vítimas. A cada notícia "
  "acrescentada a uma ocorrência, todas as notícias daquela ocorrência são "
  "relidas em conjunto e a lista de pessoas atingidas é reconstruída do zero. O "
  "procedimento é mais custoso do que processar apenas o texto novo, porém é o "
  "que impede a soma indevida: uma pessoa descrita em três fontes diferentes "
  "deve ser contada uma vez, e o total nunca pode ser inferior ao maior número "
  "citado em uma única fonte.")
p("Correções feitas pelo analista prevalecem sobre a saída automática nos "
  "reprocessamentos seguintes. Quando o modelo passa a discordar de um registro "
  "já corrigido — por exemplo, ao identificar em notícia posterior que uma "
  "pessoa ferida veio a falecer —, a divergência é apresentada como proposta "
  "sujeita a aceitação, preservando a correção humana até decisão explícita.")

quebra()

# ----------------------------------------------------------- reprodutibilidade
h(1, "11 Reprodutibilidade")

h(2, "11.1 Execução em unidade gráfica ou em processador")
p("O treino executa indiferentemente em unidade de processamento gráfico ou em "
  "processador comum. A seleção é automática: a opção padrão emprega a unidade "
  "gráfica quando disponível e recai sobre o processador quando não há uma. A "
  "escolha pode ser forçada por parâmetro de linha de comando, o que permite "
  "reproduzir o trabalho em equipamento sem placa dedicada. Em processador, o "
  "tempo de treino é substancialmente maior, mas o procedimento e os resultados "
  "são os mesmos.")
p("A distinção entre os dois modos reside apenas na variante da biblioteca de "
  "tensores instalada. A instalação a partir do arquivo de dependências traz a "
  "variante para processador, que funciona em qualquer equipamento; a variante "
  "com suporte à unidade gráfica é instalada em comando separado, substituindo "
  "a anterior:")
codigo([
    "# instalação base (variante para processador)",
    "pip install -r requirements.txt",
    "python -m spacy download pt_core_news_sm",
    "",
    "# opcional: variante com suporte a unidade gráfica NVIDIA",
    "pip install torch --index-url https://download.pytorch.org/whl/cu126",
    "",
    "# retorno à variante para processador",
    "pip install torch --index-url https://download.pytorch.org/whl/cpu",
])
p("O ambiente efetivamente utilizado em cada execução — versão da biblioteca, "
  "disponibilidade de unidade gráfica, modelo da placa e dispositivo "
  "selecionado — é registrado junto às métricas e exibido na aplicação, de modo "
  "que o relato de tempo de treino seja interpretável.")

h(2, "11.2 Execução dos experimentos")
p("A comparação completa é disparada por linha de comando, com as variações "
  "abaixo:")
codigo([
    "# sorteia a divisão e compara os seis modelos na tarefa de escopo",
    "python rodar_treino.py --dividir --tarefa violencia",
    "",
    "# executa as três tarefas em sequência",
    "python rodar_treino.py --tarefa todas",
    "",
    "# força execução em processador",
    "python rodar_treino.py --tarefa violencia --dispositivo cpu",
    "",
    "# restringe a um modelo e a uma configuração da grade",
    "python rodar_treino.py --tarefa motivacao --modelo bertimbau-base --grade 1",
])

h(2, "11.3 Registro dos resultados")
p("Cada execução grava uma linha em banco de dados contendo tarefa, modelo, "
  "hiperparâmetros, métricas completas, matriz de confusão, histórico por "
  "época, tamanho das três partições, duração e dispositivo. Os pesos do melhor "
  "modelo de cada tarefa são gravados em disco e utilizados pela interface de "
  "teste. O corpus anotado é exportável em planilha, para anexação ao trabalho "
  "e verificação independente.")

h(2, "11.4 Determinismo")
p("A semente do sorteio da divisão é explícita e registrada. A inicialização da "
  "camada de classificação e a ordem de apresentação dos lotes permanecem "
  "sujeitas a variação entre execuções, de modo que repetições da mesma "
  "configuração produzem valores próximos, porém não idênticos. Diferenças "
  "pequenas entre modelos devem, por isso, ser interpretadas com cautela, e a "
  "repetição da execução vencedora com sementes distintas é recomendada antes de "
  "afirmar superioridade.")

quebra()

# --------------------------------------------------------------- resultados
h(1, "12 Resultados")
p("Esta seção é preenchida com a saída das execuções conduzidas sobre o corpus "
  "anotado. A aplicação produz as tabelas abaixo automaticamente, e os valores "
  "devem ser transcritos a partir dela, acompanhados da data de execução e do "
  "tamanho do corpus naquele momento.")
p("Enquanto o corpus estiver em construção, a seção permanece sem números: "
  "métricas calculadas sobre poucos exemplos por classe variam de forma ampla "
  "com o acerto ou erro de um único caso e não sustentam conclusão, conforme "
  "discutido na seção 4.5.")

h(2, "12.1 Caracterização do corpus")
p("Reportar o total de notícias anotadas, a distribuição entre as classes da "
  "tarefa de escopo, a distribuição por categoria de motivação e por indicador, "
  "e o tamanho das três partições. Classes com suporte reduzido devem ser "
  "identificadas nesta subseção, pois condicionam a leitura das métricas.")
p()
tabela(
    ["Partição", "Total", "Com disparo", "Sem disparo"],
    [["Treino", "", "", ""], ["Validação", "", "", ""],
     ["Teste", "", "", ""], ["Total", "", "", ""]],
    [4.0, 4.0, 4.0, 4.0],
)
legenda("Tabela 6 — Composição do corpus por partição (a preencher).")

h(2, "12.2 Tarefa de escopo")
p("Reportar, para cada modelo e configuração, as métricas e a matriz de "
  "confusão. Destacar a configuração vencedora segundo o critério da seção 8.2 e "
  "discutir a distribuição dos dois tipos de erro à luz da assimetria "
  "apresentada na seção 9.2.")
p()
tabela(
    ["Modelo", "Taxa", "Lote", "Épocas", "Acurácia", "Precisão", "Recall", "F1 macro"],
    [[m["rotulo"], "", "", "", "", "", "", ""] for m in mod_treino.MODELOS.values()],
    [3.6, 1.6, 1.2, 1.6, 2.0, 1.8, 1.6, 2.0],
)
legenda("Tabela 7 — Desempenho na tarefa de escopo, melhor configuração por modelo "
        "(a preencher).")

h(2, "12.3 Tarefa de motivação")
p("Reportar as mesmas métricas e a matriz de confusão completa das "
  f"{len(CATEGORIAS)} categorias. A análise deve privilegiar os pares de "
  "categorias com maior confusão mútua, verificando se correspondem a fronteiras "
  "reconhecidamente difíceis do manual de curadoria — hipótese que, se "
  "confirmada, indica revisão de critério de anotação em vez de limitação do "
  "modelo.")

h(2, "12.4 Tarefa de indicadores")
p("Reportar as médias macro e micro, as métricas por indicador e a matriz 2×2 "
  "de cada um, conforme a seção 2.6.")

h(2, "12.5 Contribuição da camada de modelo de linguagem")
p("Comparar a saída do classificador com a decisão final após a validação, "
  "quantificando em que proporção dos casos a camada confirma, corrige ou "
  "reverte a classificação. A comparação é viabilizada pela interface descrita "
  "na seção 9.3 e responde se a camada custosa agrega decisão própria ou apenas "
  "reproduz o estágio anterior.")

quebra()

# --------------------------------------------------------------- limitacoes
h(1, "13 Limitações e ameaças à validade")

h(2, "13.1 Validade interna")
p("A anotação foi conduzida por um único anotador, o que impede o cálculo de "
  "concordância entre anotadores e deixa sem estimativa a parcela de erro "
  "atribuível à subjetividade do rótulo. A anotação dupla de uma amostra do "
  "corpus, com cálculo de coeficiente de concordância, é a mitigação indicada.")
p("A busca de hiperparâmetros percorre grade reduzida, selecionada por "
  "conhecimento prévio da literatura, e não uma exploração ampla do espaço. O "
  "desempenho reportado é, portanto, um limite inferior do alcançável.")

h(2, "13.2 Validade externa")
p("O corpus reflete os veículos de imprensa e as regiões efetivamente cobertas "
  "durante a sua construção. Modelos treinados sobre ele podem degradar diante "
  "de vocabulário regional distinto ou de convenções de redação de outros "
  "veículos. A avaliação em conjunto proveniente de fontes não representadas no "
  "treino é o teste adequado para essa limitação.")

h(2, "13.3 Viés de cobertura")
p("Indicadores construídos a partir de notícias herdam o viés de noticiabilidade "
  "da imprensa. Episódios em determinadas localidades ou envolvendo determinados "
  "perfis de vítima são noticiados com frequência desigual, e nenhum "
  "aprimoramento de classificador corrige essa distorção de origem. O sistema "
  "mede violência armada noticiada, e os resultados devem ser enunciados nesses "
  "termos.")

h(2, "13.4 Dependência de serviço externo")
p("A camada de validação depende de serviço de terceiros, cujo modelo pode ser "
  "atualizado ou descontinuado. Como a saída do classificador é registrada "
  "separadamente da decisão final, é possível reavaliar o efeito de uma troca de "
  "modelo sem reprocessar o corpus, o que delimita o impacto dessa dependência.")

quebra()

# ---------------------------------------------------- consideracoes finais
h(1, "14 Considerações finais")
p("A documentação apresentou o método completo de um sistema de classificação "
  "de notícias de violência armada em dois estágios, com validação posterior por "
  "modelo de linguagem. As decisões de projeto foram justificadas pelos seus "
  "efeitos verificáveis: a separação entre escopo e categoria pela "
  "interpretabilidade das matrizes de confusão e pela redução de custo; a "
  "fixação da divisão amostral pela comparabilidade entre modelos; a seleção de "
  "hiperparâmetros no conjunto de validação pela validade da estimativa de "
  "desempenho; e o registro da fonte de cada exemplo pela auditabilidade do "
  "resultado.")
p("A contribuição metodológica central não está em um modelo específico, e sim "
  "no protocolo que permite afirmar algo sobre ele. O trabalho descreve um "
  "caminho verificável: um corpus cuja procedência é rastreável, uma divisão "
  "amostral estável, uma comparação ampla de arquiteturas sob busca de "
  "hiperparâmetros, métricas que não escondem classes raras e uma interface que "
  "expõe a decisão de cada camada isoladamente.")
p("Como continuidade, destacam-se a anotação dupla com medida de concordância, a "
  "ampliação do corpus nas categorias de menor suporte, a avaliação em fontes "
  "não representadas no treino e a investigação de técnicas de balanceamento "
  "para as categorias raras.")

quebra()

# -------------------------------------------------------------- referencias
h(1, "Referências")
for ref in [
    "DEVLIN, J. et al. BERT: pre-training of deep bidirectional transformers for "
    "language understanding. In: Proceedings of NAACL-HLT, 2019.",
    "SOUZA, F.; NOGUEIRA, R.; LOTUFO, R. BERTimbau: pretrained BERT models for "
    "Brazilian Portuguese. In: Brazilian Conference on Intelligent Systems "
    "(BRACIS), 2020.",
    "CONNEAU, A. et al. Unsupervised cross-lingual representation learning at "
    "scale. In: Proceedings of ACL, 2020.",
    "SANH, V. et al. DistilBERT, a distilled version of BERT: smaller, faster, "
    "cheaper and lighter. arXiv:1910.01108, 2019.",
    "RODRIGUES, J. et al. Advancing neural encoding of Portuguese with "
    "transformer Albertina PT-*. arXiv:2305.06721, 2023.",
    "VASWANI, A. et al. Attention is all you need. In: Advances in Neural "
    "Information Processing Systems, 2017.",
    "SOKOLOVA, M.; LAPALME, G. A systematic analysis of performance measures for "
    "classification tasks. Information Processing & Management, v. 45, n. 4, 2009.",
    "BRASIL. Lei nº 8.069, de 13 de julho de 1990. Estatuto da Criança e do "
    "Adolescente.",
]:
    par = doc.add_paragraph()
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.space_after = Pt(10)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.add_run(ref)

quebra()

# ----------------------------------------------------------------- apendices
h(1, "Apêndice A — Categorias de motivação")
p("Reprodução das definições utilizadas na anotação e no prompt do modelo de "
  "linguagem. A coluna de desambiguação registra as fronteiras que mais geram "
  "divergência.")
p()
tabela(
    ["Categoria", "Definição", "Desambiguação"],
    [[c["nome"], c["definicao"], c["nao_e"]] for c in CATEGORIAS],
    [3.0, 6.5, 6.5],
)
legenda(f"Tabela 8 — As {len(CATEGORIAS)} categorias de motivação.")

quebra()

h(1, "Apêndice B — Indicadores")
p("Os indicadores dividem-se em dois tipos. Os de regra são calculados sobre a "
  "lista consolidada de vítimas e recalculados a cada alteração; os de texto "
  "resultam da classificação do conteúdo e são confirmados pela camada de modelo "
  "de linguagem.")
p()
tabela(
    ["Indicador", "Tipo", "Critério"],
    [[i["rotulo"], "regra" if i["tipo"] == "regra" else "texto", i["definicao"]]
     for i in INDICADORES],
    [4.2, 1.8, 10.0],
)
legenda("Tabela 9 — Indicadores e respectivos critérios.")

p("Os indicadores de regra são determinísticos e não dependem de modelo: "
  "vítima com menos de dezoito anos, vítima do gênero feminino, vítima agente "
  "de segurança, vítima com cargo ou candidatura política, e três ou mais "
  "mortos na mesma ocorrência. Por serem calculados sobre a lista final, "
  "atualizam-se sozinhos quando a situação de uma vítima muda — uma pessoa "
  "ferida que vem a falecer pode converter a ocorrência em chacina.")
p("Os indicadores de texto são os que exigem classificação e, portanto, os que "
  "entram na avaliação da segunda etapa. O cadastro de novos indicadores de "
  "texto exige exemplos anotados para retreinamento, requisito incorporado ao "
  "fluxo da aplicação de gestão.")

quebra()

h(1, "Apêndice C — Diretrizes de anotação")
p("Resumo operacional para quem conduz a anotação.")
p()
tabela(
    ["Situação", "Decisão"],
    [
        ["Notícia relata tiro efetuado, com ou sem vítima atingida", "Escopo: sim"],
        ["Ameaça com arma de fogo, sem disparo", "Escopo: não"],
        ["Apreensão de arma ou munição, sem confronto", "Escopo: não"],
        ["Morte ou lesão por arma branca ou outro meio", "Escopo: não"],
        ["Bala perdida proveniente de confronto", "Escopo: sim; motivação segue o confronto de origem"],
        ["Disparo acidental no manuseio da arma", "Escopo: sim; motivação disparo_acidental"],
        ["Tiro para o alto, sem alvo", "Escopo: sim; motivação tiros_a_esmo"],
        ["Polícia presente e atirando no momento do fato", "Motivação acao_operacao_policial"],
        ["Polícia chega após o fato, apenas para atender", "Não caracteriza ação policial"],
        ["Assalto em que se atira para intimidar", "Motivação roubo_tentativa, não ataque_a_civis"],
        ["Morte como desfecho de outro motivo relatado", "O outro motivo é o principal"],
        ["Texto insuficiente para decidir a motivação", "Rotular apenas o escopo; registrar a dúvida na observação"],
    ],
    [8.0, 8.0,],
)
legenda("Tabela 10 — Diretrizes para casos de fronteira.")

p("Duas recomendações gerais. Primeiro, a motivação a registrar é a do fato que "
  "originou os tiros, e não a do seu desfecho nem de acontecimentos anteriores "
  "ou posteriores sem relação com o disparo. Segundo, em caso de dúvida "
  "genuína, rotula-se o escopo e deixa-se a motivação em branco, registrando a "
  "razão no campo de observação: um rótulo de motivação incerto prejudica mais a "
  "segunda tarefa do que a sua ausência.")

caminho = SAIDA / "TCC-Classificador-Violencia-Armada.docx"
doc.save(caminho)
print("gerado:", caminho)
print("tamanho:", round(caminho.stat().st_size / 1024), "KB")
print("paragrafos:", len(doc.paragraphs), "| tabelas:", len(doc.tables))
